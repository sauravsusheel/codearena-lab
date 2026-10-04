import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BASE = r"c:\Users\bdeepak\OneDrive - Chanakya University\Course materials\projects"
DARK = RGBColor(0x1A,0x23,0x3A)
BLUE = RGBColor(0x1F,0x50,0x9E)
GRAY = RGBColor(0x60,0x60,0x60)

def shd(cell,c):
    tc=cell._tc; p=tc.get_or_add_tcPr(); s=OxmlElement("w:shd")
    s.set(qn("w:val"),"clear"); s.set(qn("w:color"),"auto"); s.set(qn("w:fill"),c); p.append(s)

def bdr(t):
    for r in t._tbl.findall(qn("w:tr")):
        for c in r.findall(qn("w:tc")):
            p=c.get_or_add_tcPr(); b=OxmlElement("w:tcBorders")
            for s in("top","left","bottom","right","insideH","insideV"):
                e=OxmlElement("w:"+s); e.set(qn("w:val"),"single"); e.set(qn("w:sz"),"4")
                e.set(qn("w:space"),"0"); e.set(qn("w:color"),"BFCDE0"); b.append(e)
            p.append(b)

def hd(d,t,lv=1):
    p=d.add_heading(t,level=lv); r=p.runs[0] if p.runs else p.add_run(t)
    r.font.color.rgb=DARK if lv==1 else BLUE; r.font.bold=True
    r.font.size=Pt(16 if lv==1 else 13 if lv==2 else 11)
    p.paragraph_format.space_before=Pt(12); p.paragraph_format.space_after=Pt(4)

def bd(d,t):
    p=d.add_paragraph(); r=p.add_run(t); r.font.size=Pt(10.5); r.font.color.rgb=GRAY
    p.paragraph_format.space_after=Pt(4)

def nt(d,t):
    p=d.add_paragraph(); p.paragraph_format.left_indent=Inches(0.25)
    p.paragraph_format.space_before=Pt(6); p.paragraph_format.space_after=Pt(6)
    r1=p.add_run("Note:  "); r1.font.bold=True; r1.font.size=Pt(10); r1.font.color.rgb=BLUE
    r2=p.add_run(t); r2.font.size=Pt(10); r2.font.italic=True; r2.font.color.rgb=GRAY
    pp=p._p.get_or_add_pPr(); pb=OxmlElement("w:pBdr"); lf=OxmlElement("w:left")
    lf.set(qn("w:val"),"single"); lf.set(qn("w:sz"),"12"); lf.set(qn("w:space"),"4"); lf.set(qn("w:color"),"1F509E")
    pb.append(lf); pp.append(pb)

def cb(d,t):
    p=d.add_paragraph(); p.paragraph_format.left_indent=Inches(0.25); p.paragraph_format.right_indent=Inches(0.25)
    p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(8)
    r=p.add_run(t); r.font.name="Courier New"; r.font.size=Pt(8.5); r.font.color.rgb=RGBColor(0x2D,0x2D,0x2D)
    pp=p._p.get_or_add_pPr(); s=OxmlElement("w:shd")
    s.set(qn("w:val"),"clear"); s.set(qn("w:color"),"auto"); s.set(qn("w:fill"),"F0F4FA"); pp.append(s)

def bl(d,t,px=None):
    p=d.add_paragraph(style="List Bullet")
    if px:
        r=p.add_run(px+"  "); r.font.bold=True; r.font.size=Pt(10.5); r.font.color.rgb=BLUE
    r2=p.add_run(t); r2.font.size=Pt(10.5); r2.font.color.rgb=GRAY; p.paragraph_format.space_after=Pt(2)

def tb(d,hs,rs,ws=None):
    t=d.add_table(rows=1+len(rs),cols=len(hs)); t.alignment=WD_TABLE_ALIGNMENT.LEFT; t.style="Table Grid"
    hc=t.rows[0].cells
    for i,h in enumerate(hs):
        hc[i].text=h; hc[i].paragraphs[0].runs[0].font.bold=True; hc[i].paragraphs[0].runs[0].font.size=Pt(10)
        hc[i].paragraphs[0].runs[0].font.color.rgb=RGBColor(0xFF,0xFF,0xFF); shd(hc[i],"1F509E")
        hc[i].vertical_alignment=WD_ALIGN_VERTICAL.CENTER
    for ri,row in enumerate(rs):
        cs=t.rows[ri+1].cells
        for ci,v in enumerate(row):
            cs[ci].text=str(v); cs[ci].paragraphs[0].runs[0].font.size=Pt(10)
            cs[ci].paragraphs[0].runs[0].font.color.rgb=GRAY
            if ri%2==0: shd(cs[ci],"F5F8FF")
    bdr(t)
    if ws:
        for ro in t.rows:
            for ci,c in enumerate(ro.cells): c.width=Inches(ws[ci])
    d.add_paragraph(); return t

def nd():
    doc=Document()
    for s in doc.sections:
        s.top_margin=Cm(2.5); s.bottom_margin=Cm(2.5); s.left_margin=Cm(3.0); s.right_margin=Cm(2.5)
    doc.styles["Normal"].font.name="Calibri"; doc.styles["Normal"].font.size=Pt(10.5); return doc

def cv(doc,no,sh,fu,po,db,dp):
    doc.add_paragraph()
    for tx,sz,bo,co in [("CHANAKYA UNIVERSITY",14,True,DARK),("Department of Computer Science and Applications",11,False,GRAY)]:
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(tx)
        r.font.size=Pt(sz); r.font.bold=bo; r.font.color.rgb=co
    doc.add_paragraph()
    for tx,sz,bo,co in [(f"UniSphere Academic Portal  |  Module {no:02d}",13,True,BLUE),(fu,19,True,DARK),(f"Team Dossier  --  {sh}",11,False,GRAY)]:
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(tx)
        r.font.size=Pt(sz); r.font.bold=bo; r.font.color.rgb=co
    doc.add_paragraph()
    tb(doc,["Parameter","Value"],[
        ["Team Number",f"Team {no:02d}"],["Module Name",fu],
        ["Programme","Master of Computer Applications (MCA) -- Semester 3"],
        ["Academic Year","2025-2026"],["Service Port",str(po)],
        ["Database Name",db],["Database Port",str(dp)],
        ["Document Version","1.0"],["Classification","Internal Academic Use Only"]
    ],ws=[2.0,4.0])
    doc.add_page_break()

RH=["Role","Responsibility","Sprint Ownership"]; RW=[2.2,3.0,1.8]
SH=["Sprint","Weeks","Deliverables"]; SW=[1.0,0.8,5.2]
AH=["Method","Endpoint","Required Role","Description"]
