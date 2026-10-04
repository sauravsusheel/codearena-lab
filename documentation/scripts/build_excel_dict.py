import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = r"c:\Users\bdeepak\OneDrive - Chanakya University\Course materials\projects"
DOCS_DIR = os.path.join(BASE, "docs")
os.makedirs(DOCS_DIR, exist_ok=True)

wb = openpyxl.Workbook()
wb.remove(wb.active)

NAVY_FILL = PatternFill(start_color="002561", end_color="002561", fill_type="solid")
SAFFRON_FILL = PatternFill(start_color="DB6C29", end_color="DB6C29", fill_type="solid")
HEADER_SUB_FILL = PatternFill(start_color="1F509E", end_color="1F509E", fill_type="solid")
ZEBRA_FILL = PatternFill(start_color="F8F9FA", end_color="F8F9FA", fill_type="solid")
WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

PK_FILL = PatternFill(start_color="E6EEF8", end_color="E6EEF8", fill_type="solid")
FK_FILL = PatternFill(start_color="FEF3E9", end_color="FEF3E9", fill_type="solid")
UQ_FILL = PatternFill(start_color="FCF6E5", end_color="FCF6E5", fill_type="solid")

FONT_TITLE = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
FONT_HEADER = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
FONT_REGULAR = Font(name="Calibri", size=10, color="212529")
FONT_REGULAR_BOLD = Font(name="Calibri", size=10, bold=True, color="212529")
FONT_CODE = Font(name="Consolas", size=9.5, color="1A233A")
FONT_PK = Font(name="Calibri", size=10, bold=True, color="002561")
FONT_FK = Font(name="Calibri", size=10, bold=True, color="9C410A")
FONT_UQ = Font(name="Calibri", size=10, bold=True, color="855D08")

BORDER_THIN = Border(
    left=Side(style="thin", color="DEE2E6"), right=Side(style="thin", color="DEE2E6"),
    top=Side(style="thin", color="DEE2E6"), bottom=Side(style="thin", color="DEE2E6")
)
BORDER_HEADER = Border(
    left=Side(style="thin", color="002561"), right=Side(style="thin", color="002561"),
    top=Side(style="thin", color="002561"), bottom=Side(style="medium", color="002561")
)

def apply_key_styling(cell, key_type):
    if "PRIMARY" in key_type and "FOREIGN" in key_type:
        cell.fill = PK_FILL; cell.font = FONT_PK
    elif "PRIMARY" in key_type:
        cell.fill = PK_FILL; cell.font = FONT_PK
    elif "FOREIGN" in key_type:
        cell.fill = FK_FILL; cell.font = FONT_FK
    elif "UNIQUE" in key_type:
        cell.fill = UQ_FILL; cell.font = FONT_UQ
    else:
        cell.font = FONT_REGULAR

MODULES = []
MODULES.append({
    "code": "T01", "name": "EduHRMS", "title": "Human Resource Management System",
    "db": "hrms_db", "port": 5432, "sheet_name": "T01_HRMS",
    "tables": [
        {"name": "hrms_users", "desc": "Faculty and staff user directory with authentication and roles",
         "columns": [
             ("id", "UUID", "PRIMARY KEY", "NO", "gen_random_uuid()", "None", "Primary key user ID"),
             ("email", "VARCHAR(255)", "UNIQUE KEY", "NO", "None", "None", "Official institutional email"),
             ("password_hash", "VARCHAR(255)", "NONE", "NO", "None", "None", "Argon2id password hash"),
             ("full_name", "VARCHAR(128)", "NONE", "NO", "None", "None", "Full legal name"),
             ("role", "VARCHAR(32)", "NONE", "NO", "'STAFF'", "None", "SUPER_ADMIN, DEAN, HOD, FACULTY, STAFF"),
             ("dept_code", "VARCHAR(16)", "NONE", "NO", "None", "None", "Department code (e.g. MCA, CSE)"),
             ("dept_name", "VARCHAR(128)", "NONE", "NO", "None", "None", "Department title"),
             ("is_active", "BOOLEAN", "NONE", "NO", "TRUE", "None", "Active employment status"),
             ("created_at", "TIMESTAMPTZ", "NONE", "NO", "NOW()", "None", "Record creation timestamp"),
             ("updated_at", "TIMESTAMPTZ", "NONE", "NO", "NOW()", "None", "Last update timestamp")
         ]},
        {"name": "hrms_faculty_profiles", "desc": "Faculty tenure, academic rank, salary and credit workload",
         "columns": [
             ("id", "UUID", "PRIMARY KEY", "NO", "gen_random_uuid()", "None", "Profile record unique ID"),
             ("user_id", "UUID", "FOREIGN KEY", "NO", "None", "hrms_users(id)", "References hrms_users(id)"),
             ("designation", "VARCHAR(64)", "NONE", "NO", "None", "None", "Rank (e.g., Associate Professor)"),
             ("tenure_status", "VARCHAR(32)", "NONE", "NO", "'PROBATION'", "None", "PROBATION, CONFIRMED, SENIOR, PROFESSOR"),
             ("date_of_joining", "DATE", "NONE", "NO", "None", "None", "Date joined university"),
             ("base_salary", "NUMERIC(12,2)", "NONE", "NO", "None", "None", "Basic monthly salary"),
             ("max_teaching_credits", "INT", "NONE", "NO", "16", "None", "Maximum semester teaching credits"),
             ("phone", "VARCHAR(16)", "NONE", "YES", "None", "None", "Contact phone number"),
             ("profile_photo_url", "VARCHAR(512)", "NONE", "YES", "None", "None", "Profile image URL in MinIO"),
             ("created_at", "TIMESTAMPTZ", "NONE", "NO", "NOW()", "None", "Profile creation timestamp")
         ]},
        {"name": "hrms_leave_requests", "desc": "Digital leave workflow with multi-tier approvals",
         "columns": [
             ("id", "UUID", "PRIMARY KEY", "NO", "gen_random_uuid()", "None", "Leave request unique ID"),
             ("user_id", "UUID", "FOREIGN KEY", "NO", "None", "hrms_users(id)", "Applicant user ID"),
             ("leave_type", "VARCHAR(32)", "NONE", "NO", "None", "None", "Casual (CL), Medical (ML), Duty (DL)"),
             ("start_date", "DATE", "NONE", "NO", "None", "None", "Leave starting date"),
             ("end_date", "DATE", "NONE", "NO", "None", "None", "Leave conclusion date"),
             ("days_requested", "INT", "NONE", "NO", "None", "None", "Total calendar days"),
             ("reason", "TEXT", "NONE", "NO", "None", "None", "Detailed justification"),
             ("status", "VARCHAR(32)", "NONE", "NO", "'PENDING_HOD'", "None", "PENDING_HOD, PENDING_DEAN, APPROVED, REJECTED"),
             ("hod_action_at", "TIMESTAMPTZ", "NONE", "YES", "None", "None", "HOD action timestamp"),
             ("dean_action_at", "TIMESTAMPTZ", "NONE", "YES", "None", "None", "Dean action timestamp"),
             ("approved_by", "UUID", "FOREIGN KEY", "YES", "None", "hrms_users(id)", "Final approving authority user ID"),
             ("rejection_reason", "TEXT", "NONE", "YES", "None", "None", "Mandatory rejection comment"),
             ("created_at", "TIMESTAMPTZ", "NONE", "NO", "NOW()", "None", "Submission timestamp")
         ]},
        {"name": "hrms_payroll_records", "desc": "Automated monthly payroll, allowances, deductions, and payslips",
         "columns": [
             ("id", "UUID", "PRIMARY KEY", "NO", "gen_random_uuid()", "None", "Payroll record unique ID"),
             ("user_id", "UUID", "FOREIGN KEY", "NO", "None", "hrms_users(id)", "Employee user ID"),
             ("month_year", "VARCHAR(16)", "NONE", "NO", "None", "None", "Billing cycle YYYY-MM (Unique with user_id)"),
             ("basic_pay", "NUMERIC(12,2)", "NONE", "NO", "None", "None", "Basic monthly salary"),
             ("hra", "NUMERIC(10,2)", "NONE", "NO", "0.00", "None", "House Rent Allowance"),
             ("da", "NUMERIC(10,2)", "NONE", "NO", "0.00", "None", "Dearness Allowance"),
             ("pf_deduction", "NUMERIC(10,2)", "NONE", "NO", "0.00", "None", "Provident fund deduction"),
             ("tds", "NUMERIC(10,2)", "NONE", "NO", "0.00", "None", "Tax Deducted at Source"),
             ("net_pay", "NUMERIC(12,2)", "NONE", "NO", "None", "None", "Net payable salary"),
             ("pdf_payslip_url", "VARCHAR(512)", "NONE", "YES", "None", "None", "Pre-signed S3 download URL"),
             ("created_at", "TIMESTAMPTZ", "NONE", "NO", "NOW()", "None", "Generation timestamp")
         ]},
        {"name": "hrms_audit_logs", "desc": "Immutable security and compliance log for HR transactions",
         "columns": [
             ("id", "UUID", "PRIMARY KEY", "NO", "gen_random_uuid()", "None", "Audit log unique ID"),
             ("actor_id", "UUID", "FOREIGN KEY", "NO", "None", "hrms_users(id)", "User initiating transaction"),
             ("action", "VARCHAR(64)", "NONE", "NO", "None", "None", "Action code (e.g. LEAVE_APPROVED)"),
             ("table_name", "VARCHAR(64)", "NONE", "NO", "None", "None", "Target database table"),
             ("record_id", "UUID", "NONE", "YES", "None", "None", "Affected entity ID"),
             ("payload_diff", "JSONB", "NONE", "YES", "None", "None", "Field state diff JSON"),
             ("ip_address", "VARCHAR(45)", "NONE", "YES", "None", "None", "Client IP address"),
             ("timestamp", "TIMESTAMPTZ", "NONE", "NO", "NOW()", "None", "Event timestamp")
         ]}
    ]
})
MODULES.append({
    "code": "T02", "name": "SmartAttend", "title": "Smart Student Attendance Tracking System",
    "db": "attend_db", "port": 5433, "sheet_name": "T02_Attendance",
    "tables": [
        {"name": "attend_courses", "desc": "Course master catalog linked to teaching faculty and semester cohorts",
         "columns": [
             ("id", "UUID", "PRIMARY KEY", "NO", "gen_random_uuid()", "None", "Course record unique ID"),
             ("course_code", "VARCHAR(16)", "UNIQUE KEY", "NO", "None", "None", "Course code (e.g., MCA301)"),
             ("course_name", "VARCHAR(128)", "NONE", "NO", "None", "None", "Course full academic title"),
             ("batch", "VARCHAR(32)", "NONE", "NO", "None", "None", "Cohort batch (e.g. 2024-2026)"),
             ("semester", "VARCHAR(32)", "NONE", "NO", "None", "None", "Semester (e.g., Semester 3)"),
             ("faculty_id", "UUID", "NONE", "NO", "None", "None", "Assigned faculty user_id from Gateway"),
             ("max_students", "INT", "NONE", "NO", "60", "None", "Class enrollment ceiling"),
             ("created_at", "TIMESTAMPTZ", "NONE", "NO", "NOW()", "None", "Creation timestamp")
         ]},
        {"name": "attend_sessions", "desc": "Class sessions with dynamic QR codes and geofence radius",
         "columns": [
             ("id", "UUID", "PRIMARY KEY", "NO", "gen_random_uuid()", "None", "Session unique ID"),
             ("course_id", "UUID", "FOREIGN KEY", "NO", "None", "attend_courses(id)", "References attend_courses(id)"),
             ("session_date", "DATE", "NONE", "NO", "None", "None", "Lecture date"),
             ("start_time", "TIME", "NONE", "NO", "None", "None", "Scheduled start time"),
             ("end_time", "TIME", "NONE", "NO", "None", "None", "Scheduled end time"),
             ("room_code", "VARCHAR(16)", "NONE", "YES", "None", "None", "Classroom/Lab number"),
             ("geo_lat", "NUMERIC(9,6)", "NONE", "YES", "None", "None", "Classroom GPS latitude"),
             ("geo_lng", "NUMERIC(9,6)", "NONE", "YES", "None", "None", "Classroom GPS longitude"),
             ("geo_radius_m", "INT", "NONE", "NO", "50", "None", "Geofence tolerance radius in metres"),
             ("qr_token_hash", "VARCHAR(64)", "NONE", "YES", "None", "None", "SHA-256 hash of active QR code"),
             ("qr_token_expiry", "TIMESTAMPTZ", "NONE", "YES", "None", "None", "Token expiry (90 seconds lifetime)"),
             ("status", "VARCHAR(16)", "NONE", "NO", "'SCHEDULED'", "None", "SCHEDULED, OPEN, CLOSED, CANCELLED"),
             ("opened_by", "UUID", "NONE", "YES", "None", "None", "Faculty user_id opening session"),
             ("opened_at", "TIMESTAMPTZ", "NONE", "YES", "None", "None", "Session start timestamp"),
             ("closed_at", "TIMESTAMPTZ", "NONE", "YES", "None", "None", "Session closure timestamp"),
             ("created_at", "TIMESTAMPTZ", "NONE", "NO", "NOW()", "None", "Creation timestamp")
         ]},
        {"name": "attend_records", "desc": "Student attendance logs with verification coordinates and mode",
         "columns": [
             ("id", "UUID", "PRIMARY KEY", "NO", "gen_random_uuid()", "None", "Attendance record ID"),
             ("session_id", "UUID", "FOREIGN KEY", "NO", "None", "attend_sessions(id)", "References attend_sessions(id)"),
             ("student_id", "UUID", "NONE", "NO", "None", "None", "Student user_id (Unique with session_id)"),
             ("student_roll_no", "VARCHAR(32)", "NONE", "NO", "None", "None", "Student roll number"),
             ("status", "VARCHAR(16)", "NONE", "NO", "'ABSENT'", "None", "PRESENT, ABSENT, LATE"),
             ("mode", "VARCHAR(16)", "NONE", "NO", "None", "None", "QR, GPS, MANUAL, REGULARIZED"),
             ("device_lat", "NUMERIC(9,6)", "NONE", "YES", "None", "None", "Device GPS latitude"),
             ("device_lng", "NUMERIC(9,6)", "NONE", "YES", "None", "None", "Device GPS longitude"),
             ("marked_at", "TIMESTAMPTZ", "NONE", "NO", "NOW()", "None", "Verification timestamp")
         ]},
        {"name": "attend_regularisation", "desc": "Absence regularisation requests with medical/emergency proof",
         "columns": [
             ("id", "UUID", "PRIMARY KEY", "NO", "gen_random_uuid()", "None", "Regularisation request ID"),
             ("student_id", "UUID", "NONE", "NO", "None", "None", "Student user_id requesting relief"),
             ("session_id", "UUID", "FOREIGN KEY", "NO", "None", "attend_sessions(id)", "Target missed class session"),
             ("reason", "TEXT", "NONE", "NO", "None", "None", "Absence justification"),
             ("evidence_url", "VARCHAR(512)", "NONE", "YES", "None", "None", "Medical certificate PDF URL"),
             ("status", "VARCHAR(16)", "NONE", "NO", "'PENDING'", "None", "PENDING, APPROVED, REJECTED"),
             ("reviewed_by", "UUID", "NONE", "YES", "None", "None", "Faculty reviewer user_id"),
             ("review_comment", "TEXT", "NONE", "YES", "None", "None", "Faculty review notes"),
             ("created_at", "TIMESTAMPTZ", "NONE", "NO", "NOW()", "None", "Submission timestamp")
         ]}
    ]
})

MODULES.append({
    "code": "T03", "name": "ExamShield", "title": "AI-Powered Secure Online Examination System",
    "db": "exam_db", "port": 5434, "sheet_name": "T03_Exams",
    "tables": [
        {"name": "exam_questions", "desc": "Question bank tagged with Bloom taxonomy, rubrics, and difficulty",
         "columns": [
             ("id", "UUID", "PRIMARY KEY", "NO", "gen_random_uuid()", "None", "Question unique identifier"),
             ("course_code", "VARCHAR(16)", "NONE", "NO", "None", "None", "Course curriculum code"),
             ("question_text", "TEXT", "NONE", "NO", "None", "None", "Problem statement / question body"),
             ("question_type", "VARCHAR(16)", "NONE", "NO", "None", "None", "MCQ, TRUE_FALSE, SHORT, DESCRIPTIVE, CODE"),
             ("difficulty", "VARCHAR(16)", "NONE", "NO", "'MEDIUM'", "None", "EASY, MEDIUM, HARD"),
             ("bloom_level", "VARCHAR(32)", "NONE", "YES", "None", "None", "Remember, Understand, Apply, Analyze, Evaluate"),
             ("marks", "INT", "NONE", "NO", "1", "None", "Question maximum marks"),
             ("options", "JSONB", "NONE", "YES", "None", "None", "MCQ options and answer keys JSON"),
             ("model_answer", "TEXT", "NONE", "YES", "None", "None", "Model answer for AI grading"),
             ("grading_rubric", "TEXT", "NONE", "YES", "None", "None", "Rubric guidelines for Gemini AI"),
             ("tags", "TEXT[]", "NONE", "YES", "None", "None", "Topic taxonomy tags"),
             ("created_by", "UUID", "NONE", "NO", "None", "None", "Faculty author user_id"),
             ("created_at", "TIMESTAMPTZ", "NONE", "NO", "NOW()", "None", "Creation timestamp")
         ]},
        {"name": "exam_papers", "desc": "Scheduled exams with duration, window dates, and pass standards",
         "columns": [
             ("id", "UUID", "PRIMARY KEY", "NO", "gen_random_uuid()", "None", "Exam paper unique ID"),
             ("title", "VARCHAR(256)", "NONE", "NO", "None", "None", "Examination title"),
             ("course_code", "VARCHAR(16)", "NONE", "NO", "None", "None", "Associated course code"),
             ("duration_mins", "INT", "NONE", "NO", "None", "None", "Test duration in minutes"),
             ("total_marks", "INT", "NONE", "NO", "None", "None", "Total maximum score"),
             ("pass_marks", "INT", "NONE", "NO", "None", "None", "Passing threshold marks"),
             ("scheduled_start", "TIMESTAMPTZ", "NONE", "NO", "None", "None", "Window opening time"),
             ("scheduled_end", "TIMESTAMPTZ", "NONE", "NO", "None", "None", "Window closing time"),
             ("status", "VARCHAR(16)", "NONE", "NO", "'DRAFT'", "None", "DRAFT, SCHEDULED, LIVE, COMPLETED, ARCHIVED"),
             ("allow_retakes", "BOOLEAN", "NONE", "NO", "FALSE", "None", "Retake permission flag"),
             ("created_by", "UUID", "NONE", "NO", "None", "None", "Faculty author user_id"),
             ("created_at", "TIMESTAMPTZ", "NONE", "NO", "NOW()", "None", "Creation timestamp")
         ]},
        {"name": "exam_attempts", "desc": "Student test attempts tracking proctoring flags and score",
         "columns": [
             ("id", "UUID", "PRIMARY KEY", "NO", "gen_random_uuid()", "None", "Attempt unique ID"),
             ("paper_id", "UUID", "FOREIGN KEY", "NO", "None", "exam_papers(id)", "References exam_papers(id) (Unique with student_id)"),
             ("student_id", "UUID", "NONE", "NO", "None", "None", "Candidate student user_id"),
             ("started_at", "TIMESTAMPTZ", "NONE", "NO", "NOW()", "None", "Exam start timestamp"),
             ("submitted_at", "TIMESTAMPTZ", "NONE", "YES", "None", "None", "Exam submit timestamp"),
             ("status", "VARCHAR(16)", "NONE", "NO", "'NOT_STARTED'", "None", "NOT_STARTED, IN_PROGRESS, SUBMITTED, GRADED"),
             ("tab_switch_count", "INT", "NONE", "NO", "0", "None", "Browser focus loss infractions"),
             ("proctoring_flags", "JSONB", "NONE", "NO", "'[]'", "None", "Logged suspicious AI events"),
             ("total_score", "NUMERIC(6,2)", "NONE", "YES", "None", "None", "Final aggregate score")
         ]},
        {"name": "exam_responses", "desc": "Question answers evaluated by auto-grader, Gemini AI, and faculty",
         "columns": [
             ("id", "UUID", "PRIMARY KEY", "NO", "gen_random_uuid()", "None", "Response unique ID"),
             ("attempt_id", "UUID", "FOREIGN KEY", "NO", "None", "exam_attempts(id)", "References exam_attempts(id)"),
             ("question_id", "UUID", "NONE", "NO", "None", "None", "Target question ID"),
             ("response_text", "TEXT", "NONE", "YES", "None", "None", "Submitted written response"),
             ("code_submission", "TEXT", "NONE", "YES", "None", "None", "Submitted coding solution"),
             ("language", "VARCHAR(16)", "NONE", "YES", "None", "None", "Programming language used"),
             ("auto_score", "NUMERIC(5,2)", "NONE", "YES", "None", "None", "Objective auto-grade"),
             ("ai_score", "NUMERIC(5,2)", "NONE", "YES", "None", "None", "Gemini AI subjective rubric grade"),
             ("final_score", "NUMERIC(5,2)", "NONE", "YES", "None", "None", "Final recorded mark"),
             ("faculty_override", "BOOLEAN", "NONE", "NO", "FALSE", "None", "True if human examiner changed score"),
             ("feedback", "TEXT", "NONE", "YES", "None", "None", "Qualitative critique"),
             ("graded_at", "TIMESTAMPTZ", "NONE", "YES", "None", "None", "Evaluation timestamp")
         ]}
    ]
})
MODULES.append({
    "code": "T04", "name": "CodeArena", "title": "Coding Practice and Assignment Platform",
    "db": "code_db", "port": 5435, "sheet_name": "T04_CodeArena",
    "tables": [
        {"name": "code_problems", "desc": "Algorithmic problem catalogue with limits, constraints, and editorials",
         "columns": [
             ("id", "UUID", "PRIMARY KEY", "NO", "gen_random_uuid()", "None", "Problem unique identifier"),
             ("title", "VARCHAR(256)", "NONE", "NO", "None", "None", "Problem display title"),
             ("slug", "VARCHAR(256)", "UNIQUE KEY", "NO", "None", "None", "URL slug identifier"),
             ("difficulty", "VARCHAR(16)", "NONE", "NO", "'MEDIUM'", "None", "EASY, MEDIUM, HARD"),
             ("description", "TEXT", "NONE", "NO", "None", "None", "Problem statement"),
             ("constraints", "TEXT", "NONE", "YES", "None", "None", "Input mathematical bounds"),
             ("sample_input", "TEXT", "NONE", "YES", "None", "None", "Sample input case"),
             ("sample_output", "TEXT", "NONE", "YES", "None", "None", "Sample expected output"),
             ("editorial", "TEXT", "NONE", "YES", "None", "None", "Reference solution guide"),
             ("time_limit_ms", "INT", "NONE", "NO", "2000", "None", "CPU limit in ms"),
             ("memory_limit_kb", "INT", "NONE", "NO", "262144", "None", "RAM limit in KB (256 MB)"),
             ("tags", "TEXT[]", "NONE", "YES", "None", "None", "Taxonomy tags (DP, Graphs, Trees)"),
             ("created_by", "UUID", "NONE", "NO", "None", "None", "Faculty author user_id"),
             ("is_public", "BOOLEAN", "NONE", "NO", "TRUE", "None", "Public catalog visibility"),
             ("created_at", "TIMESTAMPTZ", "NONE", "NO", "NOW()", "None", "Creation timestamp")
         ]},
        {"name": "code_test_cases", "desc": "Input/output test suites used by Judge0 CE sandbox",
         "columns": [
             ("id", "UUID", "PRIMARY KEY", "NO", "gen_random_uuid()", "None", "Test case unique ID"),
             ("problem_id", "UUID", "FOREIGN KEY", "NO", "None", "code_problems(id)", "References code_problems(id)"),
             ("input", "TEXT", "NONE", "NO", "None", "None", "Standard input stream content"),
             ("expected", "TEXT", "NONE", "NO", "None", "None", "Expected stdout stream"),
             ("is_sample", "BOOLEAN", "NONE", "NO", "FALSE", "None", "True if visible sample case"),
             ("weight", "INT", "NONE", "NO", "1", "None", "Score weight")
         ]},
        {"name": "code_submissions", "desc": "Candidate code executions, verdicts, execution metrics, and AI reviews",
         "columns": [
             ("id", "UUID", "PRIMARY KEY", "NO", "gen_random_uuid()", "None", "Submission unique ID"),
             ("problem_id", "UUID", "FOREIGN KEY", "NO", "None", "code_problems(id)", "Target problem ID"),
             ("student_id", "UUID", "NONE", "NO", "None", "None", "Student user_id"),
             ("language", "VARCHAR(16)", "NONE", "NO", "None", "None", "python, cpp, java, c, go, javascript"),
             ("source_code", "TEXT", "NONE", "NO", "None", "None", "Submitted source code"),
             ("verdict", "VARCHAR(16)", "NONE", "NO", "'PENDING'", "None", "AC, WA, TLE, MLE, RE, CE, PENDING"),
             ("runtime_ms", "INT", "NONE", "YES", "None", "None", "Execution time in ms"),
             ("memory_kb", "INT", "NONE", "YES", "None", "None", "RAM consumed in KB"),
             ("score", "NUMERIC(5,2)", "NONE", "NO", "0.00", "None", "Test suite pass percentage"),
             ("ai_feedback", "TEXT", "NONE", "YES", "None", "None", "Gemini AI code critique"),
             ("submitted_at", "TIMESTAMPTZ", "NONE", "NO", "NOW()", "None", "Submission timestamp")
         ]},
        {"name": "code_contests", "desc": "Timed competitive programming events with ICPC scoring",
         "columns": [
             ("id", "UUID", "PRIMARY KEY", "NO", "gen_random_uuid()", "None", "Contest unique ID"),
             ("title", "VARCHAR(256)", "NONE", "NO", "None", "None", "Contest title"),
             ("starts_at", "TIMESTAMPTZ", "NONE", "NO", "None", "None", "Contest start time"),
             ("ends_at", "TIMESTAMPTZ", "NONE", "NO", "None", "None", "Contest end time"),
             ("created_by", "UUID", "NONE", "NO", "None", "None", "Faculty author user_id"),
             ("created_at", "TIMESTAMPTZ", "NONE", "NO", "NOW()", "None", "Creation timestamp")
         ]}
    ]
})

MODULES.append({
    "code": "T05", "name": "PresentAI", "title": "AI-Evaluated Online Presentation and Viva System",
    "db": "pres_db", "port": 5436, "sheet_name": "T05_PresentationAI",
    "tables": [
        {"name": "pres_sessions", "desc": "Presentation sessions scheduled with course associations and rubrics",
         "columns": [
             ("id", "UUID", "PRIMARY KEY", "NO", "gen_random_uuid()", "None", "Session unique ID"),
             ("course_code", "VARCHAR(16)", "NONE", "NO", "None", "None", "Course curriculum code"),
             ("faculty_id", "UUID", "NONE", "NO", "None", "None", "Faculty examiner user_id"),
             ("title", "VARCHAR(256)", "NONE", "NO", "None", "None", "Session topic headline"),
             ("scheduled_at", "TIMESTAMPTZ", "NONE", "NO", "None", "None", "Scheduled start time"),
             ("duration_mins", "INT", "NONE", "NO", "15", "None", "Duration per slot in minutes"),
             ("state", "VARCHAR(16)", "NONE", "NO", "'SCHEDULED'", "None", "SCHEDULED, WAITING, LIVE, PROCESSING, COMPLETED"),
             ("created_at", "TIMESTAMPTZ", "NONE", "NO", "NOW()", "None", "Creation timestamp")
         ]},
        {"name": "pres_slots", "desc": "Student presentation slots with slide uploads, speech transcripts, and recordings",
         "columns": [
             ("id", "UUID", "PRIMARY KEY", "NO", "gen_random_uuid()", "None", "Slot unique ID"),
             ("session_id", "UUID", "FOREIGN KEY", "NO", "None", "pres_sessions(id)", "References session (Unique with student_id)"),
             ("student_id", "UUID", "NONE", "NO", "None", "None", "Presenting student user_id"),
             ("slide_url", "VARCHAR(512)", "NONE", "YES", "None", "None", "Uploaded slide deck link (PDF/PPTX)"),
             ("recording_url", "VARCHAR(512)", "NONE", "YES", "None", "None", "MinIO stored video recording URL"),
             ("transcript", "TEXT", "NONE", "YES", "None", "None", "Speech-to-Text verbatim transcript"),
             ("scheduled_start", "TIMESTAMPTZ", "NONE", "NO", "None", "None", "Slot start time"),
             ("started_at", "TIMESTAMPTZ", "NONE", "YES", "None", "None", "Commencement timestamp"),
             ("ended_at", "TIMESTAMPTZ", "NONE", "YES", "None", "None", "Conclusion timestamp")
         ]},
        {"name": "pres_rubric_criteria", "desc": "Rubric dimensions (Content, Delivery, Q&A, Time) and score caps",
         "columns": [
             ("id", "UUID", "PRIMARY KEY", "NO", "gen_random_uuid()", "None", "Criterion unique ID"),
             ("session_id", "UUID", "FOREIGN KEY", "NO", "None", "pres_sessions(id)", "References pres_sessions(id)"),
             ("criterion_name", "VARCHAR(64)", "NONE", "NO", "None", "None", "Dimension (e.g. Content, Delivery)"),
             ("max_score", "INT", "NONE", "NO", "None", "None", "Maximum marks allowed"),
             ("description", "TEXT", "NONE", "YES", "None", "None", "Assessment criteria guidelines")
         ]},
        {"name": "pres_evaluations", "desc": "Per-criterion scoring storing AI ratings and faculty examiner marks",
         "columns": [
             ("id", "UUID", "PRIMARY KEY", "NO", "gen_random_uuid()", "None", "Evaluation entry ID"),
             ("slot_id", "UUID", "FOREIGN KEY", "NO", "None", "pres_slots(id)", "References pres_slots(id) (Unique with criterion_id)"),
             ("criterion_id", "UUID", "FOREIGN KEY", "NO", "None", "pres_rubric_criteria(id)", "References rubric criterion"),
             ("ai_score", "NUMERIC(5,2)", "NONE", "YES", "None", "None", "Gemini AI automated score"),
             ("ai_justification", "TEXT", "NONE", "YES", "None", "None", "AI scoring rationale"),
             ("faculty_score", "NUMERIC(5,2)", "NONE", "YES", "None", "None", "Faculty assigned score"),
             ("faculty_comment", "TEXT", "NONE", "YES", "None", "None", "Examiner feedback notes"),
             ("final_score", "NUMERIC(5,2)", "NONE", "YES", "None", "None", "Final official mark"),
             ("status", "VARCHAR(16)", "NONE", "NO", "'PENDING'", "None", "PENDING, AI_DONE, FACULTY_REVIEWED, FINALISED")
         ]}
    ]
})
MODULES.append({
    "code": "T06", "name": "CampusBot", "title": "RAG-Powered Campus AI Assistant",
    "db": "bot_db", "port": 5437, "sheet_name": "T06_CampusBot",
    "tables": [
        {"name": "bot_documents", "desc": "Ingested university circulars, syllabi, handbooks, and knowledge documents",
         "columns": [
             ("id", "UUID", "PRIMARY KEY", "NO", "gen_random_uuid()", "None", "Document record unique ID"),
             ("title", "VARCHAR(256)", "NONE", "NO", "None", "None", "Document title"),
             ("source_type", "VARCHAR(32)", "NONE", "NO", "None", "None", "PDF, DOCX, TXT, CIRCULAR"),
             ("category", "VARCHAR(64)", "NONE", "YES", "None", "None", "Admissions, Exams, Fees, Syllabus"),
             ("uploaded_by", "UUID", "NONE", "NO", "None", "None", "Admin uploader user_id"),
             ("version", "INT", "NONE", "NO", "1", "None", "Document version counter"),
             ("is_active", "BOOLEAN", "NONE", "NO", "TRUE", "None", "Active knowledge retrieval flag"),
             ("created_at", "TIMESTAMPTZ", "NONE", "NO", "NOW()", "None", "Ingestion timestamp")
         ]},
        {"name": "bot_chunks", "desc": "512-token segments with 768-dim vector embeddings for cosine similarity",
         "columns": [
             ("id", "UUID", "PRIMARY KEY", "NO", "gen_random_uuid()", "None", "Chunk unique ID"),
             ("document_id", "UUID", "FOREIGN KEY", "NO", "None", "bot_documents(id)", "References parent document"),
             ("chunk_index", "INT", "NONE", "NO", "None", "None", "Sequential index in document"),
             ("chunk_text", "TEXT", "NONE", "NO", "None", "None", "Segment text content"),
             ("embedding", "vector(768)", "NONE", "YES", "None", "None", "Google text-embedding-004 vector"),
             ("token_count", "INT", "NONE", "YES", "None", "None", "Token length of chunk"),
             ("created_at", "TIMESTAMPTZ", "NONE", "NO", "NOW()", "None", "Embedding timestamp")
         ]},
        {"name": "bot_conversations", "desc": "User dialogue sessions maintaining rolling multi-turn context",
         "columns": [
             ("id", "UUID", "PRIMARY KEY", "NO", "gen_random_uuid()", "None", "Conversation thread unique ID"),
             ("user_id", "UUID", "NONE", "NO", "None", "None", "Student or faculty user_id"),
             ("started_at", "TIMESTAMPTZ", "NONE", "NO", "NOW()", "None", "Thread commencement timestamp"),
             ("ended_at", "TIMESTAMPTZ", "NONE", "YES", "None", "None", "Thread termination timestamp")
         ]},
        {"name": "bot_messages", "desc": "Individual chat messages with source chunk attribution and feedback",
         "columns": [
             ("id", "UUID", "PRIMARY KEY", "NO", "gen_random_uuid()", "None", "Message unique ID"),
             ("conversation_id", "UUID", "FOREIGN KEY", "NO", "None", "bot_conversations(id)", "References conversation"),
             ("role", "VARCHAR(16)", "NONE", "NO", "None", "None", "user, assistant, system"),
             ("content", "TEXT", "NONE", "NO", "None", "None", "Message text content"),
             ("source_chunks", "UUID[]", "NONE", "YES", "None", "None", "Array of chunk IDs referenced"),
             ("feedback", "SMALLINT", "NONE", "YES", "None", "None", "User rating: +1 or -1"),
             ("created_at", "TIMESTAMPTZ", "NONE", "NO", "NOW()", "None", "Message dispatch timestamp")
         ]}
    ]
})

MODULES.append({
    "code": "T07", "name": "UniHub", "title": "University Research and Project Repository Portal",
    "db": "hub_db", "port": 5438, "sheet_name": "T07_UniHub",
    "tables": [
        {"name": "hub_projects", "desc": "University repository index synced with self-hosted Gitea CE",
         "columns": [
             ("id", "UUID", "PRIMARY KEY", "NO", "gen_random_uuid()", "None", "Project unique identifier"),
             ("title", "VARCHAR(256)", "NONE", "NO", "None", "None", "Project title"),
             ("description", "TEXT", "NONE", "YES", "None", "None", "Project synopsis"),
             ("department", "VARCHAR(64)", "NONE", "NO", "None", "None", "Host academic department"),
             ("academic_year", "VARCHAR(16)", "NONE", "NO", "None", "None", "Academic year cohort"),
             ("tech_stack", "TEXT[]", "NONE", "YES", "None", "None", "Frameworks and languages used"),
             ("status", "VARCHAR(16)", "NONE", "NO", "'IN_PROGRESS'", "None", "IDEATION, IN_PROGRESS, COMPLETED, PUBLISHED"),
             ("gitea_repo_url", "VARCHAR(512)", "NONE", "YES", "None", "None", "Gitea clone URL"),
             ("gitea_repo_id", "INT", "NONE", "YES", "None", "None", "Gitea repository numeric ID"),
             ("star_count", "INT", "NONE", "NO", "0", "None", "Bookmark / star count"),
             ("is_public", "BOOLEAN", "NONE", "NO", "TRUE", "None", "Public showcase visibility"),
             ("created_by", "UUID", "NONE", "NO", "None", "None", "Author student or faculty user_id"),
             ("created_at", "TIMESTAMPTZ", "NONE", "NO", "NOW()", "None", "Registration timestamp")
         ]},
        {"name": "hub_project_members", "desc": "Team member associations and contributor roles",
         "columns": [
             ("project_id", "UUID", "PRIMARY & FOREIGN KEY", "NO", "None", "hub_projects(id)", "References hub_projects(id)"),
             ("user_id", "UUID", "PRIMARY KEY", "NO", "None", "None", "Team member user_id"),
             ("role", "VARCHAR(32)", "NONE", "NO", "'CONTRIBUTOR'", "None", "LEAD, CONTRIBUTOR, MENTOR"),
             ("joined_at", "TIMESTAMPTZ", "NONE", "NO", "NOW()", "None", "Team association timestamp")
         ]},
        {"name": "hub_papers", "desc": "Research publications and preprints with full-text trigram indexing",
         "columns": [
             ("id", "UUID", "PRIMARY KEY", "NO", "gen_random_uuid()", "None", "Paper unique identifier"),
             ("title", "VARCHAR(512)", "NONE", "NO", "None", "None", "Paper full title"),
             ("abstract", "TEXT", "NONE", "NO", "None", "None", "Paper abstract (indexed via pg_trgm)"),
             ("keywords", "TEXT[]", "NONE", "YES", "None", "None", "Research keyword tags"),
             ("doi", "VARCHAR(128)", "NONE", "YES", "None", "None", "Digital Object Identifier"),
             ("journal_name", "VARCHAR(256)", "NONE", "YES", "None", "None", "Journal or conference name"),
             ("published_year", "SMALLINT", "NONE", "YES", "None", "None", "Publication year"),
             ("pdf_url", "VARCHAR(512)", "NONE", "YES", "None", "None", "MinIO PDF URL"),
             ("visibility", "VARCHAR(16)", "NONE", "NO", "'PUBLIC'", "None", "DRAFT, PUBLIC, RESTRICTED"),
             ("citation_count", "INT", "NONE", "NO", "0", "None", "Institutional citations count"),
             ("uploaded_by", "UUID", "NONE", "NO", "None", "None", "Author or uploader user_id"),
             ("created_at", "TIMESTAMPTZ", "NONE", "NO", "NOW()", "None", "Upload timestamp")
         ]},
        {"name": "hub_activity_feed", "desc": "University-wide real-time activity stream of code pushes and papers",
         "columns": [
             ("id", "UUID", "PRIMARY KEY", "NO", "gen_random_uuid()", "None", "Activity log unique ID"),
             ("actor_id", "UUID", "NONE", "NO", "None", "None", "Acting user_id"),
             ("event_type", "VARCHAR(32)", "NONE", "NO", "None", "None", "REPO_PUSH, PAPER_UPLOADED, MILESTONE_REACHED"),
             ("entity_id", "UUID", "NONE", "YES", "None", "None", "Target entity ID"),
             ("entity_type", "VARCHAR(32)", "NONE", "YES", "None", "None", "PROJECT, PAPER, MILESTONE"),
             ("summary", "TEXT", "NONE", "NO", "None", "None", "Human readable text description"),
             ("created_at", "TIMESTAMPTZ", "NONE", "NO", "NOW()", "None", "Event timestamp")
         ]}
    ]
})
MODULES.append({
    "code": "T08", "name": "TrackIt", "title": "Project and Internship Track Management System",
    "db": "track_db", "port": 5439, "sheet_name": "T08_TrackIt",
    "tables": [
        {"name": "track_projects", "desc": "Academic project registrations, faculty guide allocation, and lifecycle stages",
         "columns": [
             ("id", "UUID", "PRIMARY KEY", "NO", "gen_random_uuid()", "None", "Project record unique identifier"),
             ("title", "VARCHAR(256)", "NONE", "NO", "None", "None", "Project title"),
             ("abstract", "TEXT", "NONE", "NO", "None", "None", "Project abstract"),
             ("objectives", "TEXT", "NONE", "YES", "None", "None", "Functional deliverable goals"),
             ("domain", "VARCHAR(64)", "NONE", "NO", "None", "None", "Domain (AI, Cloud, Cyber, IoT)"),
             ("tech_stack", "TEXT[]", "NONE", "YES", "None", "None", "Proposed technology stack"),
             ("project_type", "VARCHAR(16)", "NONE", "NO", "None", "None", "MINI, MAJOR, RESEARCH, OPEN_SOURCE"),
             ("stage", "VARCHAR(16)", "NONE", "NO", "'PROPOSED'", "None", "PROPOSED, APPROVED, IN_PROGRESS, COMPLETED"),
             ("lead_student_id", "UUID", "NONE", "NO", "None", "None", "Lead student user_id"),
             ("guide_faculty_id", "UUID", "NONE", "YES", "None", "None", "Assigned faculty guide user_id"),
             ("academic_year", "VARCHAR(16)", "NONE", "NO", "None", "None", "Academic year cohort"),
             ("semester", "VARCHAR(32)", "NONE", "NO", "None", "None", "Semester designation"),
             ("unihub_repo_id", "UUID", "NONE", "YES", "None", "None", "Linked UniHub repository ID"),
             ("created_at", "TIMESTAMPTZ", "NONE", "NO", "NOW()", "None", "Registration timestamp"),
             ("updated_at", "TIMESTAMPTZ", "NONE", "NO", "NOW()", "None", "Last modification timestamp")
         ]},
        {"name": "track_project_members", "desc": "Student team membership and role designations for group projects",
         "columns": [
             ("project_id", "UUID", "PRIMARY & FOREIGN KEY", "NO", "None", "track_projects(id)", "References track_projects(id)"),
             ("student_id", "UUID", "PRIMARY KEY", "NO", "None", "None", "Team member student user_id"),
             ("role", "VARCHAR(32)", "NONE", "NO", "'MEMBER'", "None", "Role: LEAD or MEMBER"),
             ("joined_at", "TIMESTAMPTZ", "NONE", "NO", "NOW()", "None", "Team association timestamp")
         ]},
        {"name": "track_milestones", "desc": "Milestone checkpoints configured by faculty guide with due dates",
         "columns": [
             ("id", "UUID", "PRIMARY KEY", "NO", "gen_random_uuid()", "None", "Milestone unique identifier"),
             ("project_id", "UUID", "FOREIGN KEY", "NO", "None", "track_projects(id)", "References track_projects(id)"),
             ("title", "VARCHAR(128)", "NONE", "NO", "None", "None", "Milestone title"),
             ("description", "TEXT", "NONE", "YES", "None", "None", "Deliverable requirements"),
             ("due_date", "DATE", "NONE", "NO", "None", "None", "Completion deadline"),
             ("weightage", "NUMERIC(4,1)", "NONE", "NO", "10.0", "None", "Percentage of total project mark"),
             ("status", "VARCHAR(16)", "NONE", "NO", "'PENDING'", "None", "PENDING, SUBMITTED, REVIEWED, MISSED"),
             ("order_index", "INT", "NONE", "NO", "None", "None", "Sequential index order (1 to 8)")
         ]},
        {"name": "track_milestone_submissions", "desc": "Student milestone deliverables, uploaded code/reports, and reviews",
         "columns": [
             ("id", "UUID", "PRIMARY KEY", "NO", "gen_random_uuid()", "None", "Submission unique ID"),
             ("milestone_id", "UUID", "FOREIGN KEY", "NO", "None", "track_milestones(id)", "References track_milestones(id)"),
             ("submitted_by", "UUID", "NONE", "NO", "None", "None", "Student user_id submitting"),
             ("progress_text", "TEXT", "NONE", "NO", "None", "None", "Progress report summary"),
             ("deliverable_url", "VARCHAR(512)", "NONE", "YES", "None", "None", "MinIO or GitHub deliverable link"),
             ("submitted_at", "TIMESTAMPTZ", "NONE", "NO", "NOW()", "None", "Submission timestamp"),
             ("score", "NUMERIC(5,2)", "NONE", "YES", "None", "None", "Faculty assigned milestone score"),
             ("faculty_feedback", "TEXT", "NONE", "YES", "None", "None", "Faculty guidance critique"),
             ("reviewed_at", "TIMESTAMPTZ", "NONE", "YES", "None", "None", "Review timestamp")
         ]},
        {"name": "track_mentor_logs", "desc": "Meeting interaction diary between students and faculty guide",
         "columns": [
             ("id", "UUID", "PRIMARY KEY", "NO", "gen_random_uuid()", "None", "Log unique identifier"),
             ("project_id", "UUID", "FOREIGN KEY", "NO", "None", "track_projects(id)", "References track_projects(id)"),
             ("logged_by", "UUID", "NONE", "NO", "None", "None", "User recording meeting"),
             ("meeting_date", "DATE", "NONE", "NO", "None", "None", "Meeting session date"),
             ("duration_mins", "INT", "NONE", "YES", "None", "None", "Duration in minutes"),
             ("summary", "TEXT", "NONE", "NO", "None", "None", "Discussion summary"),
             ("action_items", "TEXT", "NONE", "YES", "None", "None", "Committed action items"),
             ("created_at", "TIMESTAMPTZ", "NONE", "NO", "NOW()", "None", "Logging timestamp")
         ]},
        {"name": "track_evaluations", "desc": "Final examination committee rubric marks (Guide, Internal, External)",
         "columns": [
             ("id", "UUID", "PRIMARY KEY", "NO", "gen_random_uuid()", "None", "Evaluation unique ID"),
             ("project_id", "UUID", "FOREIGN KEY", "NO", "None", "track_projects(id)", "References project (Unique with examiner_id)"),
             ("examiner_id", "UUID", "NONE", "NO", "None", "None", "Examiner faculty user_id"),
             ("examiner_role", "VARCHAR(32)", "NONE", "NO", "None", "None", "Role: GUIDE, INTERNAL, EXTERNAL"),
             ("report_score", "NUMERIC(5,2)", "NONE", "YES", "None", "None", "Report documentation score (max 20)"),
             ("presentation_score", "NUMERIC(5,2)", "NONE", "YES", "None", "None", "Presentation score (max 20)"),
             ("technical_score", "NUMERIC(5,2)", "NONE", "YES", "None", "None", "Technical depth score (max 20)"),
             ("innovation_score", "NUMERIC(5,2)", "NONE", "YES", "None", "None", "Innovation score (max 20)"),
             ("viva_score", "NUMERIC(5,2)", "NONE", "YES", "None", "None", "Oral viva score (max 20)"),
             ("total_score", "NUMERIC(5,2)", "NONE", "YES", "None", "None", "Aggregate score (max 100)"),
             ("remarks", "TEXT", "NONE", "YES", "None", "None", "Committee comments"),
             ("submitted_at", "TIMESTAMPTZ", "NONE", "NO", "NOW()", "None", "Submission timestamp")
         ]},
        {"name": "track_internships", "desc": "Industry internship registrations, offer letter verification, and final grades",
         "columns": [
             ("id", "UUID", "PRIMARY KEY", "NO", "gen_random_uuid()", "None", "Internship unique ID"),
             ("student_id", "UUID", "NONE", "NO", "None", "None", "Intern student user_id"),
             ("company_name", "VARCHAR(256)", "NONE", "NO", "None", "None", "Employer organization name"),
             ("role_title", "VARCHAR(128)", "NONE", "NO", "None", "None", "Intern position title"),
             ("stipend_per_month", "NUMERIC(10,2)", "NONE", "YES", "None", "None", "Monthly stipend amount"),
             ("start_date", "DATE", "NONE", "NO", "None", "None", "Start date"),
             ("end_date", "DATE", "NONE", "NO", "None", "None", "End date"),
             ("offer_letter_url", "VARCHAR(512)", "NONE", "YES", "None", "None", "Offer letter PDF URL in MinIO"),
             ("status", "VARCHAR(16)", "NONE", "NO", "'REGISTERED'", "None", "REGISTERED, OFFER_VERIFIED, ONGOING, COMPLETED"),
             ("coordinator_id", "UUID", "NONE", "YES", "None", "None", "Faculty coordinator user_id"),
             ("final_grade", "VARCHAR(16)", "NONE", "YES", "None", "None", "Final awarded grade"),
             ("created_at", "TIMESTAMPTZ", "NONE", "NO", "NOW()", "None", "Registration timestamp")
         ]},
        {"name": "track_intern_logs", "desc": "Weekly work diary recording corporate hours and deliverables",
         "columns": [
             ("id", "UUID", "PRIMARY KEY", "NO", "gen_random_uuid()", "None", "Weekly log ID"),
             ("internship_id", "UUID", "FOREIGN KEY", "NO", "None", "track_internships(id)", "References internship (Unique with week_number)"),
             ("week_number", "INT", "NONE", "NO", "None", "None", "Week number index (1 to 24)"),
             ("hours_worked", "INT", "NONE", "NO", "None", "None", "Logged corporate working hours"),
             ("tasks_completed", "TEXT", "NONE", "NO", "None", "None", "Tasks and deliverables completed"),
             ("skills_acquired", "TEXT", "NONE", "YES", "None", "None", "Technical tools practiced"),
             ("challenges", "TEXT", "NONE", "YES", "None", "None", "Difficulties encountered"),
             ("submitted_at", "TIMESTAMPTZ", "NONE", "NO", "NOW()", "None", "Submission timestamp")
         ]},
        {"name": "track_intern_evaluations", "desc": "Evaluations by company supervisor and faculty coordinator",
         "columns": [
             ("id", "UUID", "PRIMARY KEY", "NO", "gen_random_uuid()", "None", "Evaluation unique ID"),
             ("internship_id", "UUID", "FOREIGN KEY", "NO", "None", "track_internships(id)", "References internship (Unique with evaluator_type)"),
             ("evaluator_type", "VARCHAR(24)", "NONE", "NO", "None", "None", "SUPERVISOR or COORDINATOR"),
             ("attendance_score", "NUMERIC(5,2)", "NONE", "YES", "None", "None", "Attendance and punctuality mark"),
             ("log_quality_score", "NUMERIC(5,2)", "NONE", "YES", "None", "None", "Log documentation precision mark"),
             ("performance_score", "NUMERIC(5,2)", "NONE", "YES", "None", "None", "Job performance evaluation mark"),
             ("remarks", "TEXT", "NONE", "YES", "None", "None", "Assessment comments"),
             ("submitted_at", "TIMESTAMPTZ", "NONE", "NO", "NOW()", "None", "Evaluation timestamp")
         ]}
    ]
})
# ==============================================================================
# SHEET 1: Master Schema Catalog
# ==============================================================================
ws_catalog = wb.create_sheet(title="Schema_Catalog")
ws_catalog.views.sheetView[0].showGridLines = True

ws_catalog.merge_cells("A1:G1")
ws_catalog["A1"] = "CHANAKYA UNIVERSITY -- UNISPHERE ACADEMIC PLATFORM"
ws_catalog["A1"].font = FONT_TITLE
ws_catalog["A1"].fill = NAVY_FILL
ws_catalog["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws_catalog.row_dimensions[1].height = 26

ws_catalog.merge_cells("A2:G2")
ws_catalog["A2"] = "Master Database Schema Directory (PostgreSQL 16 Isolated Multi-Database Architecture)"
ws_catalog["A2"].font = Font(name="Calibri", size=10.5, italic=True, color="FFFFFF")
ws_catalog["A2"].fill = SAFFRON_FILL
ws_catalog["A2"].alignment = Alignment(horizontal="center", vertical="center")
ws_catalog.row_dimensions[2].height = 20

headers_catalog = ["Module", "Team / Module Name", "Database", "Port", "Table Name", "Columns", "Purpose & Architectural Scope"]
ws_catalog.row_dimensions[4].height = 22

for col_idx, h in enumerate(headers_catalog, 1):
    c = ws_catalog.cell(row=4, column=col_idx, value=h)
    c.font = FONT_HEADER; c.fill = NAVY_FILL; c.alignment = Alignment(horizontal="center", vertical="center"); c.border = BORDER_HEADER

cur_row = 5
for mod in MODULES:
    for tbl in mod["tables"]:
        row_data = [mod["code"], mod["name"], mod["db"], mod["port"], tbl["name"], len(tbl["columns"]), tbl["desc"]]
        ws_catalog.row_dimensions[cur_row].height = 20
        fill = ZEBRA_FILL if cur_row % 2 == 0 else WHITE_FILL
        for col_idx, val in enumerate(row_data, 1):
            cell = ws_catalog.cell(row=cur_row, column=col_idx, value=val)
            cell.font = FONT_CODE if col_idx in (1, 4, 5) else FONT_REGULAR
            cell.fill = fill; cell.border = BORDER_THIN
            if col_idx in (1, 3, 4, 6):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")
        cur_row += 1

# ==============================================================================
# SHEET 2: Complete Flat Data Dictionary
# ==============================================================================
ws_dict = wb.create_sheet(title="All_Attributes_Dictionary")
ws_dict.views.sheetView[0].showGridLines = True

ws_dict.merge_cells("A1:J1")
ws_dict["A1"] = "CHANAKYA UNIVERSITY -- COMPLETE DATABASE ATTRIBUTE DICTIONARY"
ws_dict["A1"].font = FONT_TITLE
ws_dict["A1"].fill = NAVY_FILL
ws_dict["A1"].alignment = Alignment(horizontal="center", vertical="center")
ws_dict.row_dimensions[1].height = 26

ws_dict.merge_cells("A2:J2")
ws_dict["A2"] = "Universal Field-by-Field Specification for all 38 Tables across 8 Microservice Databases"
ws_dict["A2"].font = Font(name="Calibri", size=10.5, italic=True, color="FFFFFF")
ws_dict["A2"].fill = SAFFRON_FILL
ws_dict["A2"].alignment = Alignment(horizontal="center", vertical="center")
ws_dict.row_dimensions[2].height = 20

headers_dict = ["Module", "Database", "Table Name", "Attribute / Column Name", "Data Type", "Key Constraint", "Nullable", "Default Value", "Foreign Reference", "Attribute Description & Purpose"]
ws_dict.row_dimensions[4].height = 22

for col_idx, h in enumerate(headers_dict, 1):
    c = ws_dict.cell(row=4, column=col_idx, value=h)
    c.font = FONT_HEADER; c.fill = NAVY_FILL; c.alignment = Alignment(horizontal="center", vertical="center"); c.border = BORDER_HEADER

dict_row = 5
for mod in MODULES:
    for tbl in mod["tables"]:
        for col in tbl["columns"]:
            c_name, c_type, c_key, c_null, c_def, c_ref, c_desc = col
            row_data = [mod["code"], mod["db"], tbl["name"], c_name, c_type, c_key, c_null, c_def, c_ref, c_desc]
            ws_dict.row_dimensions[dict_row].height = 19
            base_fill = ZEBRA_FILL if dict_row % 2 == 0 else WHITE_FILL
            for col_idx, val in enumerate(row_data, 1):
                cell = ws_dict.cell(row=dict_row, column=col_idx, value=val)
                cell.fill = base_fill; cell.border = BORDER_THIN; cell.font = FONT_REGULAR
                if col_idx in (1, 7):
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif col_idx in (4, 5, 8, 9):
                    cell.font = FONT_CODE; cell.alignment = Alignment(horizontal="left", vertical="center")
                elif col_idx == 6:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    apply_key_styling(cell, c_key)
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
            dict_row += 1

# ==============================================================================
# SHEETS 3 to 10: Individual Module Sheets
# ==============================================================================
for mod in MODULES:
    ws = wb.create_sheet(title=mod["sheet_name"])
    ws.views.sheetView[0].showGridLines = True
    
    ws.merge_cells("A1:G1")
    ws["A1"] = f"{mod['code']} -- {mod['name']} ({mod['title']})"
    ws["A1"].font = FONT_TITLE; ws["A1"].fill = NAVY_FILL; ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26
    
    ws.merge_cells("A2:G2")
    ws["A2"] = f"PostgreSQL Database: {mod['db']}  |  Assigned Host Port: {mod['port']}  |  Tables: {len(mod['tables'])}"
    ws["A2"].font = Font(name="Calibri", size=10, italic=True, color="FFFFFF"); ws["A2"].fill = SAFFRON_FILL
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20
    
    row_cursor = 4
    for tbl in mod["tables"]:
        ws.merge_cells(start_row=row_cursor, start_column=1, end_row=row_cursor, end_column=7)
        banner = ws.cell(row=row_cursor, column=1, value=f"TABLE: {tbl['name']} -- {tbl['desc']}")
        banner.font = Font(name="Calibri", size=10.5, bold=True, color="FFFFFF")
        banner.fill = HEADER_SUB_FILL; banner.alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[row_cursor].height = 21
        row_cursor += 1
        
        col_headers = ["Attribute / Column Name", "Data Type", "Key Constraint", "Nullable", "Default Value", "Foreign Reference", "Attribute Description"]
        ws.row_dimensions[row_cursor].height = 20
        for col_idx, h in enumerate(col_headers, 1):
            c = ws.cell(row=row_cursor, column=col_idx, value=h)
            c.font = FONT_HEADER; c.fill = NAVY_FILL; c.alignment = Alignment(horizontal="center", vertical="center"); c.border = BORDER_HEADER
        row_cursor += 1
        
        for col in tbl["columns"]:
            c_name, c_type, c_key, c_null, c_def, c_ref, c_desc = col
            ws.row_dimensions[row_cursor].height = 19
            fill = ZEBRA_FILL if row_cursor % 2 == 0 else WHITE_FILL
            row_vals = [c_name, c_type, c_key, c_null, c_def, c_ref, c_desc]
            for col_idx, val in enumerate(row_vals, 1):
                cell = ws.cell(row=row_cursor, column=col_idx, value=val)
                cell.fill = fill; cell.border = BORDER_THIN; cell.font = FONT_REGULAR
                if col_idx == 1:
                    cell.font = FONT_REGULAR_BOLD; cell.alignment = Alignment(horizontal="left", vertical="center")
                elif col_idx in (2, 5, 6):
                    cell.font = FONT_CODE; cell.alignment = Alignment(horizontal="left", vertical="center")
                elif col_idx in (3, 4):
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                    if col_idx == 3:
                        apply_key_styling(cell, c_key)
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
            row_cursor += 1
        row_cursor += 1

# Widths
for sheet in wb.worksheets:
    for col in sheet.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row in (1, 2) or (sheet.title.startswith("T0") and "TABLE:" in str(cell.value)):
                continue
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 11), 52)

ws_catalog.column_dimensions["A"].width = 10
ws_catalog.column_dimensions["B"].width = 22
ws_catalog.column_dimensions["C"].width = 15
ws_catalog.column_dimensions["D"].width = 10
ws_catalog.column_dimensions["E"].width = 28
ws_catalog.column_dimensions["F"].width = 12
ws_catalog.column_dimensions["G"].width = 55

ws_dict.column_dimensions["A"].width = 10
ws_dict.column_dimensions["B"].width = 14
ws_dict.column_dimensions["C"].width = 26
ws_dict.column_dimensions["D"].width = 26
ws_dict.column_dimensions["E"].width = 16
ws_dict.column_dimensions["F"].width = 16
ws_dict.column_dimensions["G"].width = 10
ws_dict.column_dimensions["H"].width = 18
ws_dict.column_dimensions["I"].width = 24
ws_dict.column_dimensions["J"].width = 50

out_file_docs = os.path.join(DOCS_DIR, "unisphere_database_data_dictionary.xlsx")
out_file_root = os.path.join(BASE, "unisphere_database_data_dictionary.xlsx")

wb.save(out_file_docs)
wb.save(out_file_root)
print("SUCCESS: Workbook created and saved.")
