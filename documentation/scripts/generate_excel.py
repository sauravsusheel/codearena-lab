import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import csv
import os

def generate_excel():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "University Projects Catalog"

    # Ensure grid lines are visible
    ws.views.sheetView[0].showGridLines = True

    # Exact headers specified by user
    headers = [
        "Project Title1",
        "Project Problem Statement1",
        "Project Description / Proposed Solution1",
        "Expected Features / Functionalities1",
        "Suggested Technologies / Tools1",
        "Do you want to submit another project?1"
    ]

    # Style definitions
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    header_border = Border(
        left=Side(style='medium', color='1F4E79'),
        right=Side(style='medium', color='1F4E79'),
        top=Side(style='medium', color='1F4E79'),
        bottom=Side(style='medium', color='1F4E79')
    )

    # Write headers
    ws.row_dimensions[1].height = 35
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
        cell.border = header_border

    projects_data = [
        [
            "Enterprise Human Resource Management & Academic Workforce Allocation System (EduHRMS)",
            "Higher education institutions face severe administrative bottlenecks due to fragmented personnel systems, manual leave paperwork, convoluted payroll structures for academic/adjunct staff, and an absence of transparent teaching workload and credit-hour accounting. Off-the-shelf corporate HRMS solutions lack support for university-specific realities such as research grant allowances, sabbatical leave rules, examination duty compensation, and hierarchical academic approvals.",
            "A secure, web-based enterprise HRMS engineered specifically for academic institutions. It delivers automated role-based departmental workflows for faculty and staff lifecycle management, hierarchical leave approval chains (Faculty -> HOD -> Dean -> Registrar), biometric/geo-fenced staff attendance tracking, automated salary slip computation with institutional deductions and tax calculations, and dynamic teaching workload balancing (lecture credits, lab hours, office hours, and committee assignments) with full audit logging and compliance reporting.",
            "1. Multi-tier Role-Based Access Control (Admin, Dean, HOD, Faculty, Staff, Teaching Assistant).\n2. Faculty & Staff Digital Dossier (qualifications, publications, tenure tracking, emergency contacts).\n3. Intelligent Leave Management with auto-calculating leave balances (Casual, Medical, Sabbatical, Duty Leave) and multi-level approval chains.\n4. Academic Workload & Subject Allocation engine with credit-hour limit warnings and conflict detection.\n5. Payroll & Compensation module supporting base pay, PF/pension deductions, research allowance, exam duty remuneration, and PDF payslip generation.\n6. Staff Attendance Logging with geo-fencing/IP validation and monthly regularization requests.\n7. Real-time Analytics Dashboard displaying department staffing levels, absenteeism trends, and payroll expenditures.",
            "Frontend: Next.js 14, TypeScript, Tailwind CSS, Shadcn/UI, Lucide Icons, Recharts\nBackend: Node.js, Express.js (TypeScript), Zod Validation, PDFKit\nDatabase: PostgreSQL 16 (Relational Schema), Prisma ORM, Redis (Session Cache)\nDevOps & Testing: Git/GitHub, Docker, ESLint, Postman / Swagger",
            "Yes"
        ],
        [
            "Smart Campus Student Attendance Tracking, Dynamic Geo-QR & Biometric Presence Engine (EduAttend AI)",
            "Paper roll-calls and manual attendance entry squander 10-15% of instructional lecture time and remain highly vulnerable to proxy check-ins, buddy punching, and fabricated attendance excuses. Traditional static RFID cards fail because students easily pass cards to peers. Moreover, universities lack an automated, tamper-evident system to calculate real-time student attendance percentages against statutory exam qualification thresholds (e.g., mandatory 75% rule) with proactive parent/student alerts.",
            "A high-security, intelligent multi-modal student attendance tracking and eligibility management platform. The system combines time-synchronized rotating Dynamic QR codes displayed on classroom projector screens (regenerated every 5 seconds with cryptographic HMAC nonces to eliminate screenshot sharing) with student mobile GPS geofencing and Bluetooth Low Energy (BLE) beacon verification. It includes tablet-based AI facial recognition check-in with liveness detection, automated course-wise attendance tracking, an online medical leave condonation workflow, and an automated lock that directly controls examination hall ticket eligibility.",
            "1. Dynamic Rotating QR Attendance: Classroom projector displays time-expiring cryptographic QR codes refreshing every 5 seconds to eliminate screenshot sharing.\n2. Multi-Factor Geolocation & BLE Fence: Student mobile scanner verifies device latitude/longitude within classroom polygon radius and validates ambient BLE classroom beacon.\n3. AI Facial Recognition Kiosk Mode: Tablet-mounted classroom kiosk supporting high-speed multi-student facial recognition check-in with liveness detection to prevent photo spoofing.\n4. Real-time Attendance Analytics & Course Dashboards: Faculty dashboard showing instant live headcount, absent rosters, and historical attendance trends.\n5. Automated Exam Eligibility Gate: Continuous computation of attendance percentages against university threshold rules (e.g., 75% minimum, 65% with approved medical condonation), automatically locking/unlocking exam hall tickets.\n6. Medical Certificate Upload & Condonation Workflow: Student portal for uploading hospital documentation, automated OCR verification, and HOD/Dean approval state machine.\n7. Automated Defaulter Warning Escalation: Automated email and SMS notifications triggered at 85%, 80%, and <75% attendance brackets with parent notification relays.",
            "Frontend: Next.js 14, TypeScript, Tailwind CSS, Shadcn/UI, React-QR-Reader, MediaPipe (Face/Liveness)\nBackend: Node.js, Express.js (TypeScript), Socket.io (Live Headcount), Zod\nDatabase: PostgreSQL 16, Prisma ORM, Redis (5s Dynamic Nonces & Geofence Polygon Cache)\nDevOps & Testing: Git/GitHub, Docker, Jest, Postman",
            "Yes"
        ],
        [
            "AI-Proctored Intelligent Examination & Comprehensive Academic Evaluation Suite (EduExam AI)",
            "Administering remote and computer-laboratory exams in higher education suffers from critical academic integrity vulnerabilities, such as unauthorized browser tab switching, secondary device usage, impersonation, and multi-person collusion. Additionally, manual exam drafting, subjective evaluation inconsistency, and sudden student network disconnections cause severe friction and examination delays.",
            "A high-security, browser-based online examination platform with client-side AI proctoring and offline-resilient state synchronization. It enables instructors to author randomized question pools (MCQs, multi-select, and subjective essays with LaTeX/KaTeX math formula support), enforces strict client-side browser locks (disabling copy/paste, tab switching, and developer shortcuts), and integrates lightweight computer-vision models to flag unauthorized faces, eye gaze anomalies, and speech. Automated grading evaluates objective questions instantly, while a rubric-based grading console streamlines subjective assessments.",
            "1. Faculty Assessment Authoring Studio: Question banks, topic tagging, difficulty calibration, randomized question/option shuffling, and scheduled test windows.\n2. Secure Student Exam Shell: Fullscreen lock mode, disabled keyboard shortcuts (Ctrl+C, Ctrl+V, Alt+Tab, PrintScreen), and tab-switch warning counters.\n3. Real-Time AI Proctoring: In-browser WebCam/Audio monitoring detecting missing faces, multi-face presence, gaze deviations, and background speech flags.\n4. Offline Resilience & Auto-Save: Instant incremental answer caching in IndexedDB with automated WebSocket/REST sync every 15 seconds to prevent data loss upon power/internet outages.\n5. Automated Grading & Result Compilation: Instant grading for objective sections, rubric evaluation interface for subjective questions, and automated gradebook exports.\n6. Comprehensive Integrity Audit Trail: Timestamped log of browser events and webcam violation snapshots compiled into an integrity rating report for examiners.",
            "Frontend: Next.js 14, TypeScript, Tailwind CSS, TensorFlow.js / MediaPipe (Client Face/Gaze AI), KaTeX\nBackend: Node.js, Express.js (TypeScript), Socket.io (Realtime Telemetry), Zod\nDatabase: PostgreSQL 16 (JSONB for Question Schemas), Prisma ORM, Redis (Active Sessions)\nDevOps & Testing: Git/GitHub, Docker, Jest, Postman",
            "Yes"
        ],
        [
            "Cloud-Based Code Execution Arena, Automated Assessment & Plagiarism Detection Engine (EduCode Judge)",
            "Practical programming education requires frequent coding exercises and lab assessments, but manual code evaluation is slow, inconsistent, and overwhelmed by widespread code plagiarism. Furthermore, executing unverified student-submitted code on university servers introduces grave security hazards (malicious scripts, infinite resource loops, host system compromise) if not strictly isolated.",
            "A scalable, cloud-native online coding practice and automated assessment engine. It incorporates an in-browser IDE (Monaco Editor) supporting multiple languages (Python, C++, Java, JavaScript, Go), an isolated containerized execution runner with microsecond resource limits (CPU, memory, execution timeouts), automated test-case evaluation (public, private, and hidden edge cases), and an Abstract Syntax Tree (AST) code similarity and plagiarism analyzer to identify copied logic regardless of variable renaming.",
            "1. Multi-Language Web IDE: Monaco Editor integration with syntax highlighting, auto-completion, bracket matching, custom themes, and stdin input console.\n2. Secure Sandbox Execution Worker: Ephemeral Docker / gVisor execution workers with memory constraints (e.g., 256MB), CPU limits (1 core), execution timeout (3s), and network isolation.\n3. Automated Test Case Judge: Multi-tier evaluation running public test cases for immediate feedback, hidden test cases for final grading, and reporting Time Limit Exceeded (TLE), Memory Limit Exceeded (MLE), Runtime Error (RTE), and Wrong Answer (WA).\n4. Interactive Practice Arena: Curated programming problems categorized by topics (Arrays, Trees, Graphs, Dynamic Programming), user streaks, problem discussion forum, and dynamic leaderboards.\n5. Faculty Coursework Assignment Portal: Create coding assignments with starter templates, test case generators, submission deadlines, and late penalty rules.\n6. Automated Plagiarism & Code Similarity Engine: AST tokenization and Winnowing algorithm to flag identical logic structures despite variable renaming or whitespace changes.",
            "Frontend: Next.js 14, TypeScript, Tailwind CSS, Monaco Editor (@monaco-editor/react), Lucide Icons\nBackend: Node.js, Express.js, BullMQ (Distributed Task Queue), Python (AST Plagiarism Engine)\nExecution Sandbox: Docker SDK / Isolated container workers, Linux cgroups\nDatabase: PostgreSQL 16, Prisma ORM, Redis (BullMQ Queue & Session Caching)\nDevOps & Testing: Git/GitHub, Docker Compose, Jest, Postman",
            "Yes"
        ],
        [
            "AI-Powered Online Presentation, Virtual Viva Voce & Multi-Modal Evaluation Platform (EduPresent AI)",
            "Evaluating student capstone presentations, seminar defenses, and oral viva assessments remotely is highly labor-intensive, vulnerable to evaluator subjectivity, and provides students with minimal actionable feedback on delivery, vocal clarity, and visual presentation aids. Faculty examiners struggle to simultaneously assess student presentation slides, monitor body language, write notes, and formulate rigorous technical viva questions in tight 15-minute schedules.",
            "An enterprise cloud-native virtual presentation and oral defense platform equipped with multi-modal AI evaluation. It provides a real-time WebRTC virtual presentation room synchronized with slide viewing (PDF/PPTX) and an automated multi-modal evaluative pipeline. The AI engine analyzes vocal delivery (words-per-minute pace, vocal tone, clarity, and filler-word counts like 'um' and 'like'), visual engagement (eye contact, head pose, posture, and nervousness indicators), and slide semantic alignment (matching spoken words against slide bullet points using speech-to-text and LLM embeddings). The platform autonomously generates probing technical viva questions for examiners and compiles an exhaustive, rubric-scored feedback dossier.",
            "1. Live WebRTC Virtual Presentation Studio: Low-latency audio, video, and synchronized slide deck broadcast with presenter controls and audience/evaluator panels.\n2. Automated Multi-Modal Speech Analysis: Real-time audio transcription (Whisper API / Vosk), words-per-minute pace calculation, pitch variance, and filler-word counter ('uh', 'um', 'like', 'you know').\n3. Computer Vision Body Language & Eye Contact Engine: In-browser face and gaze analysis evaluating camera focus, posture stability, and nervousness metrics.\n4. Slide Content & Semantic Alignment Judge: Evaluates slide text density, layout clarity, and calculates semantic alignment between spoken audio and visual slide content via vector embeddings.\n5. AI-Assisted Viva Voce Question Generator: Real-time generation of probing, syllabus-grounded technical questions based on the student's live presentation transcript.\n6. Comprehensive Rubric-Based Feedback Dossier: Automated generation of structured evaluation reports spanning Content Mastery (30%), Oral Delivery (25%), Visual Aids (20%), Time Management (10%), and Q&A Defensibility (15%) with PDF export.\n7. Secure Recording & Institutional Archive: End-to-end encrypted session recording stored in S3/MinIO with timestamped evaluator annotations and rubrics.",
            "Frontend: Next.js 14, TypeScript, Tailwind CSS, Shadcn/UI, WebRTC (LiveKit / Mediasoup), PDF.js, MediaPipe (Face/Gaze)\nBackend: Node.js, Express.js (Signaling & Orchestration) + Python FastAPI (Multi-modal NLP & LLM Evaluator, Whisper/Vosk)\nStorage & Database: PostgreSQL 16, Prisma ORM, MinIO / S3 Object Store (Encrypted Video & Slide Storage), Redis\nDevOps & Testing: Git/GitHub, Docker, Jest, PyTest",
            "Yes"
        ],
        [
            "Institutional Knowledge Base & Context-Aware Conversational AI Assistant (EduAssist AI)",
            "University stakeholders (students, faculty, administrative staff, and visitors) are overwhelmed by fragmented PDF handbooks, scattered academic calendars, complex exam ordinances, and buried HR guidelines. Administrative departments spend excessive hours answering repetitive routine queries, while students struggle to find timely, authoritative answers regarding course requirements and academic rules.",
            "An enterprise-grade, context-aware AI conversational system powered by Retrieval-Augmented Generation (RAG). The system continuously ingests, chunks, and indexes university handbooks, syllabus documents, academic calendars, HR leave policies, and exam rules into a high-performance vector database. It provides an intuitive, real-time conversational chat interface that delivers accurate, cited answers strictly grounded in official university documents with role-aware privacy filters and a student coding tutor mode.",
            "1. Knowledge Ingestion & Document Processing Pipeline: Ingestion of PDF, DOCX, Markdown, and TXT documents with automated semantic chunking, metadata tagging, and vector embedding generation.\n2. Hybrid Semantic & Keyword Search: Vector similarity search (Cosine / HNSW) coupled with full-text keyword filtering for high-precision document retrieval.\n3. Hallucination-Guarded RAG Chatbot: Context-grounded response generation citing source documents, page numbers, and official university policies with direct link references.\n4. Role-Aware Privacy & Contextual Scoping: Dynamically adjusts response context and permissions based on user role (e.g., student queries cannot access faculty pay policy indices).\n5. AI Code & Concept Tutor Mode: Guided pedagogical assistant that provides hints and syntax explanations for university programming course queries without outputting complete assignment solutions.\n6. Administrative Analytics & Unanswered Query Escalation: Dashboard highlighting frequent student queries, knowledge gaps, user satisfaction ratings, and automated routing of unresolved questions to campus helpdesk staff.",
            "Frontend: Next.js 14, TypeScript, Tailwind CSS, Lucide Icons, React-Markdown, KaTeX\nBackend: Node.js / Express.js (Gateway & User State) + Python FastAPI (RAG Pipeline, LangChain / LlamaIndex)\nAI & Vector Store: PostgreSQL with pgvector extension (or ChromaDB), OpenAI / Gemini API / HuggingFace Embeddings\nDatabase: PostgreSQL 16, Prisma ORM, Redis (Chat Context & Rate Limiting)\nDevOps & Testing: Git/GitHub, Docker, Postman, PyTest",
            "Yes"
        ],
        [
            "Institutional Source Code Repository, Academic Research Archive & Innovation Activity Network (UniHub)",
            "Universities produce high volumes of intellectual property—including student capstone software, thesis codebases, faculty research manuscripts, and funded lab innovations. Currently, these assets remain fragmented across personal laptops, commercial Git services, and lost email threads. The institution lacks a unified code repository, a permanent academic paper pre-print archive, and a transparent discovery portal to showcase active campus research and student projects to industry partners.",
            "A self-hosted, university-wide collaborative code repository, academic research publication archive, and live campus innovation portal inspired by GitHub and arXiv. It provides Git repository hosting with web-based code browsing, branch management, commit histories, code diff viewers, and issue tracking. Concurrently, it features an institutional research repository for faculty and graduate papers with PDF viewers, BibTeX citations, and peer-review badges, paired with a real-time 'Happening Now' activity stream showcasing active project milestones, repository commits, and collaborative campus research.",
            "1. Web-Based Git Repository Hosting: Repository creation, remote Git HTTP push/pull support, file tree navigation, markdown README rendering, and syntax-highlighted source viewer.\n2. Commit History & Visual Diff Viewer: Visual commit graphs, branch switching, side-by-side / inline diffs, and lightweight issue tracking (bug reports, feature requests).\n3. University Project Showcase: Curated exhibition gallery of capstone and departmental projects with live demo links, technology tags, star ratings, and student contributor profiles.\n4. Academic Research & Paper Repository: Publication indexing with PDF preview, metadata extraction (abstract, authors, DOI, publication date), BibTeX citation generator, and departmental categorization.\n5. Campus Innovation Activity Feed ('Happening Now'): Real-time dynamic stream displaying latest commits across public university repos, new research paper uploads, project star milestones, and upcoming campus hackathons.\n6. Granular Access Control & Public/Private Visibility: Departmental, lab-level, or university-wide visibility settings ensuring confidential research remains protected while public capstones are accessible to recruiters.",
            "Frontend: Next.js 14, TypeScript, Tailwind CSS, Monaco Editor / Prism.js, react-diff-viewer, PDF.js, Lucide Icons\nBackend: Node.js, Express.js (TypeScript), isomorphic-git / nodegit, Multer (Manuscript Uploads)\nStorage & Database: PostgreSQL 16, Prisma ORM, MinIO / S3-compatible Object Store for Git Repos & PDFs, Redis\nDevOps & Testing: Git/GitHub, Docker, Jest, Postman",
            "No"
        ]
    ]

    row_fill_alt = PatternFill(start_color="F9FAFB", end_color="F9FAFB", fill_type="solid")
    row_fill_white = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    body_font = Font(name="Calibri", size=10, color="1F2937")
    body_alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    center_alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)

    for row_idx, row_data in enumerate(projects_data, 2):
        ws.row_dimensions[row_idx].height = 145
        fill = row_fill_alt if row_idx % 2 == 0 else row_fill_white
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = body_font
            cell.fill = fill
            cell.border = thin_border
            if col_idx == 6:
                cell.alignment = center_alignment
                cell.font = Font(name="Calibri", size=10, bold=True, color="0F5132" if value == "Yes" else "842029")
            else:
                cell.alignment = body_alignment

    # Column widths
    col_widths = {
        "A": 32,  # Project Title
        "B": 48,  # Problem Statement
        "C": 52,  # Description / Proposed Solution
        "D": 55,  # Expected Features
        "E": 45,  # Suggested Tech
        "F": 22   # Submit another?
    }

    for col_letter, width in col_widths.items():
        ws.column_dimensions[col_letter].width = width

    base_dir = "c:/Users/bdeepak/OneDrive - Chanakya University/Course materials/projects"
    main_file = os.path.join(base_dir, "university_projects_catalog.xlsx")
    updated_file = os.path.join(base_dir, "university_projects_catalog_updated.xlsx")
    csv_file = os.path.join(base_dir, "university_projects_catalog.csv")

    # Always save updated_file
    wb.save(updated_file)
    print(f"Successfully saved updated Excel file at {updated_file}")

    # Also save CSV
    with open(csv_file, mode='w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        for r in projects_data:
            writer.writerow(r)
    print(f"Successfully saved CSV catalog at {csv_file}")

    # Try saving to main_file if not locked
    try:
        wb.save(main_file)
        print(f"Successfully overwritten main Excel file at {main_file}")
    except PermissionError:
        print(f"Notice: '{main_file}' is currently open in Excel. Updated version is ready at '{updated_file}'.")

if __name__ == "__main__":
    generate_excel()
