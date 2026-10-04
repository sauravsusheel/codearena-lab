import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_slide_layout = prs.slide_layouts[6]

    # Color Palette
    PRIMARY_DARK = RGBColor(0x0F, 0x29, 0x4A)   # Deep Navy
    SECONDARY_BLUE = RGBColor(0x1F, 0x4E, 0x79) # Slate Navy
    ACCENT_TEAL = RGBColor(0x0D, 0x94, 0x88)    # Teal
    ACCENT_AMBER = RGBColor(0xD9, 0x77, 0x06)   # Amber
    CARD_BG = RGBColor(0xF1, 0xF5, 0xF9)        # Very light slate
    CARD_BORDER = RGBColor(0xCB, 0xD5, 0xE1)    # Light gray-blue
    TEXT_DARK = RGBColor(0x1E, 0x29, 0x3B)      # Dark text
    TEXT_MUTED = RGBColor(0x64, 0x74, 0x8B)     # Muted text
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)

    def add_header(slide, title, category="UNISPHERE ENGINEERING BLUEPRINT"):
        # Top banner category
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        tf_cat.margin_left = tf_cat.margin_right = tf_cat.margin_top = tf_cat.margin_bottom = 0
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.name = "Calibri"
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = ACCENT_TEAL

        # Title
        t_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.6))
        tf_t = t_box.text_frame
        tf_t.word_wrap = True
        tf_t.margin_left = tf_t.margin_right = tf_t.margin_top = tf_t.margin_bottom = 0
        p_t = tf_t.paragraphs[0]
        p_t.text = title
        p_t.font.name = "Calibri"
        p_t.font.size = Pt(22)
        p_t.font.bold = True
        p_t.font.color.rgb = PRIMARY_DARK

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
        return shape

    # =========================================================================
    # SLIDE 1: TITLE SLIDE
    # =========================================================================
    s1 = prs.slides.add_slide(blank_slide_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = PRIMARY_DARK
    bg1.line.fill.background()

    # Accent bar
    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.2), Inches(1.8), Inches(0.15), Inches(3.8))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ACCENT_TEAL
    bar.line.fill.background()

    t_box = s1.shapes.add_textbox(Inches(1.6), Inches(1.7), Inches(10.5), Inches(3.9))
    tf = t_box.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "UNISPHERE UNIVERSITY ECOSYSTEM"
    p0.font.name = "Calibri"
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_TEAL
    p0.space_after = Pt(8)

    p1 = tf.add_paragraph()
    p1.text = "Master Engineering Timeline, Team-Wise Workflows & Unified Integration Roadmap"
    p1.font.name = "Calibri"
    p1.font.size = Pt(28)
    p1.font.bold = True
    p1.font.color.rgb = WHITE
    p1.space_after = Pt(14)

    p2 = tf.add_paragraph()
    p2.text = "7 Specialized Modules | 21 Student Engineers (3 per Team) | 16-Week Implementation\nHRMS • Smart Attendance • AI Exams • Code Arena • AI Presentations • RAG Bot • UniHub"
    p2.font.name = "Calibri"
    p2.font.size = Pt(14)
    p2.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
    p2.space_after = Pt(24)

    p3 = tf.add_paragraph()
    p3.text = "Chanakya University • Senior Lead Architect & Cybersecurity Specification • Academic Year 2026-27"
    p3.font.name = "Calibri"
    p3.font.size = Pt(11)
    p3.font.italic = True
    p3.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # =========================================================================
    # SLIDE 2: EXECUTIVE VISION & 7 MODULAR PILLARS
    # =========================================================================
    s2 = prs.slides.add_slide(blank_slide_layout)
    add_header(s2, "Executive Vision: 7 Specialized Academic Pillars", "System Scope & Architecture")

    pillars = [
        ("Module 1: EduHRMS", "Enterprise Academic HRMS & Staff Attendance", "Faculty profiles, leave multi-approvals, payroll slips, workload balancing."),
        ("Module 2: EduAttend AI", "Smart Attendance & Exam Eligibility Gate", "5s rotating QR, GPS polygon geofence, face liveness, 75% exam gate."),
        ("Module 3: EduExam AI", "AI-Proctored Secure Online Examination", "Fullscreen lockdown, MediaPipe face/gaze telemetry, auto-sync answers."),
        ("Module 4: EduCode Judge", "Cloud Code Arena & Plagiarism Judge", "Monaco IDE, gVisor Docker sandbox, TLE/MLE judge, AST similarity."),
        ("Module 5: EduPresent AI", "AI Presentation & Virtual Viva Platform", "WebRTC stage, Whisper speech pace, eye contact, viva question generator."),
        ("Module 6: EduAssist AI", "RAG Campus Knowledge Base Assistant", "PostgreSQL pgvector, handbook/policy Q&A, cited answers, code tutor."),
        ("Module 7: UniHub", "Git Code Repository & Research Archive", "Git HTTP hosting, commit graphs, capstone showcase, pre-print papers.")
    ]

    card_w = Inches(3.7)
    card_h = Inches(2.4)
    gap_x = Inches(0.3)
    gap_y = Inches(0.25)
    start_x = Inches(0.8)
    start_y = Inches(1.5)

    for idx, (p_title, p_sub, p_desc) in enumerate(pillars[:6]):
        col = idx % 3
        row = idx // 3
        cx = start_x + col * (card_w + gap_x)
        cy = start_y + row * (card_h + gap_y)
        add_card(s2, cx, cy, card_w, card_h)

        tb = s2.shapes.add_textbox(cx + Inches(0.15), cy + Inches(0.15), card_w - Inches(0.3), card_h - Inches(0.3))
        tf_c = tb.text_frame
        tf_c.word_wrap = True
        
        p = tf_c.paragraphs[0]
        p.text = p_title
        p.font.name = "Calibri"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = SECONDARY_BLUE
        p.space_after = Pt(2)

        p_sub_p = tf_c.add_paragraph()
        p_sub_p.text = p_sub
        p_sub_p.font.name = "Calibri"
        p_sub_p.font.size = Pt(10.5)
        p_sub_p.font.bold = True
        p_sub_p.font.color.rgb = ACCENT_TEAL
        p_sub_p.space_after = Pt(4)

        p_desc_p = tf_c.add_paragraph()
        p_desc_p.text = p_desc
        p_desc_p.font.name = "Calibri"
        p_desc_p.font.size = Pt(9.5)
        p_desc_p.font.color.rgb = TEXT_DARK

    # Bottom full-width banner for Module 7 & Unified Gateway
    bot_card = add_card(s2, start_x, start_y + 2 * (card_h + gap_y) - Inches(0.1), Inches(11.7), Inches(1.0), bg_color=PRIMARY_DARK, border_color=PRIMARY_DARK)
    tb_bot = s2.shapes.add_textbox(start_x + Inches(0.2), start_y + 2 * (card_h + gap_y), Inches(11.3), Inches(0.8))
    tf_bot = tb_bot.text_frame
    tf_bot.word_wrap = True
    p_b1 = tf_bot.paragraphs[0]
    p_b1.text = "MODULE 7: UNIHUB GIT & RESEARCH ARCHIVE + UNIFIED SINGLE SIGN-ON GATEWAY"
    p_b1.font.name = "Calibri"
    p_b1.font.size = Pt(12)
    p_b1.font.bold = True
    p_b1.font.color.rgb = ACCENT_TEAL
    p_b2 = tf_bot.add_paragraph()
    p_b2.text = "Provides institutional Git hosting, research paper DOI tracking, and live campus innovation feeds, all unified by a Central JWT Auth SSO Gateway."
    p_b2.font.name = "Calibri"
    p_b2.font.size = Pt(10)
    p_b2.font.color.rgb = WHITE

    # =========================================================================
    # SLIDE 3: 16-WEEK MASTER ENGINEERING TIMELINE
    # =========================================================================
    s3 = prs.slides.add_slide(blank_slide_layout)
    add_header(s3, "16-Week Semester Engineering Timeline (6 Sprints)", "Execution Schedule")

    sprints = [
        ("Sprint 1 (Weeks 1-2)", "Inception & Schema Freeze", "• Monorepo scaffold (Turborepo)\n• Shared packages/ui setup\n• PostgreSQL BCNF schema validation\n• OpenAPI contract definition"),
        ("Sprint 2 (Weeks 3-4)", "Core Backend & Auth", "• Central JWT SSO login & roles\n• Express & FastAPI controllers\n• Database migrations deployed\n• Basic UI mockups in Next.js"),
        ("Sprint 3 (Weeks 5-6)", "Domain Core Engine", "• Leave state machine (Team 1)\n• 5s QR & Geofence (Team 2)\n• Lockdown exam shell (Team 3)\n• gVisor Docker runner (Team 4)"),
        ("Sprint 4 (Weeks 7-8)", "AI & Real-Time Telemetry", "• WebRTC presentation room (Team 5)\n• pgvector RAG pipeline (Team 6)\n• Git HTTP smart server (Team 7)\n• WebSocket state sync & alerts"),
        ("Sprint 5 (Weeks 9-10)", "Security & API Gateway", "• Reverse proxy ingress setup\n• Token blacklist in Redis\n• OWASP audit & injection hardening\n• Automated Jest test suites (>80%)"),
        ("Sprint 6 (Weeks 11-12)", "Master Merge & E2E Tests", "• Staging branch master merge\n• Attendance-to-Exam integration\n• Slide-to-Speech viva testing\n• Cross-module UI theme harmonization"),
    ]

    s_w = Inches(3.7)
    s_h = Inches(2.4)
    for idx, (s_title, s_sub, s_points) in enumerate(sprints):
        col = idx % 3
        row = idx // 3
        cx = start_x + col * (s_w + gap_x)
        cy = start_y + row * (s_h + gap_y)
        add_card(s3, cx, cy, s_w, s_h)

        tb = s3.shapes.add_textbox(cx + Inches(0.15), cy + Inches(0.15), s_w - Inches(0.3), s_h - Inches(0.3))
        tf_s = tb.text_frame
        tf_s.word_wrap = True

        p = tf_s.paragraphs[0]
        p.text = s_title
        p.font.name = "Calibri"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = SECONDARY_BLUE
        p.space_after = Pt(2)

        p_sub = tf_s.add_paragraph()
        p_sub.text = s_sub
        p_sub.font.name = "Calibri"
        p_sub.font.size = Pt(10.5)
        p_sub.font.bold = True
        p_sub.font.color.rgb = ACCENT_TEAL
        p_sub.space_after = Pt(4)

        p_pts = tf_s.add_paragraph()
        p_pts.text = s_points
        p_pts.font.name = "Calibri"
        p_pts.font.size = Pt(9.5)
        p_pts.font.color.rgb = TEXT_DARK

    bot_card3 = add_card(s3, start_x, start_y + 2 * (s_h + gap_y) - Inches(0.1), Inches(11.7), Inches(1.0), bg_color=PRIMARY_DARK, border_color=PRIMARY_DARK)
    tb_b3 = s3.shapes.add_textbox(start_x + Inches(0.2), start_y + 2 * (s_h + gap_y), Inches(11.3), Inches(0.8))
    tf_b3 = tb_b3.text_frame
    tf_b3.word_wrap = True
    p_b3_1 = tf_b3.paragraphs[0]
    p_b3_1.text = "WEEKS 13-16: LOAD TESTING, FACULTY AUDIT & CAPSTONE LIVE SHOWCASE"
    p_b3_1.font.name = "Calibri"
    p_b3_1.font.size = Pt(12)
    p_b3_1.font.bold = True
    p_b3_1.font.color.rgb = ACCENT_TEAL
    p_b3_2 = tf_b3.add_paragraph()
    p_b3_2.text = "Simulate 500 concurrent student logins, conduct live proctored coding exam, test real-time presentation viva, and present the ecosystem."
    p_b3_2.font.name = "Calibri"
    p_b3_2.font.size = Pt(10)
    p_b3_2.font.color.rgb = WHITE

    # =========================================================================
    # SLIDE 4 TO 10: TEAM-BY-TEAM WORKFLOW SLIDES
    # =========================================================================
    team_slides_data = [
        {
            "num": "Team 1",
            "title": "Module 1: Enterprise Academic HRMS (EduHRMS)",
            "cat": "Team 1 Workflow & Role Breakdown",
            "summary": "Automates faculty and staff dossier lifecycle, multi-tier leave approval state machines, teaching credit-hour allocations, and automated salary slip generation.",
            "s1": "Student 1 (Frontend & UX Lead):\n• Build faculty/staff profile dossiers & tenure tracking\n• Create multi-level leave request dialogs\n• Staff attendance check-in UI with geo-location pins\n• Departmental workforce analytics dashboard",
            "s2": "Student 2 (Backend & API Lead):\n• Implement leave state machine (Faculty -> HOD -> Dean)\n• Workload credit-hour balancing algorithm\n• Automated payroll engine with PF/tax deductions\n• Express REST controllers with Zod schema validation",
            "s3": "Student 3 (DevOps, Database & Security Lead):\n• Design normalized PostgreSQL HR & payroll tables\n• Automated PDF payslip generator pipeline (PDFKit)\n• RBAC middleware enforcement and JWT claims\n• Immutable audit log for sensitive salary adjustments"
        },
        {
            "num": "Team 2",
            "title": "Module 2: Smart Attendance & Exam Gate (EduAttend AI)",
            "cat": "Team 2 Workflow & Role Breakdown",
            "summary": "Eliminates proxy check-ins using rotating dynamic QR codes, GPS geofencing, and face liveness, while automatically gating student exam eligibility.",
            "s1": "Student 1 (Frontend & Scanner UX Lead):\n• Classroom projector dynamic rotating QR display\n• Mobile camera QR scanner with permission handling\n• Student course-wise attendance tracking cards\n• Defaulter warning banner for <75% attendance",
            "s2": "Student 2 (Backend & Geo-Fencing Architect):\n• 5-second rotating HMAC-SHA256 QR nonces in Redis\n• Geo-fence polygon point-in-polygon verification\n• Automated exam eligibility calculation service\n• WebSockets live classroom headcount broadcaster",
            "s3": "Student 3 (AI Vision & Security Specialist):\n• MediaPipe face recognition kiosk with liveness blink\n• Medical certificate OCR upload & condonation pipeline\n• Mock GPS detection (Android mock provider flags)\n• Automated parent/student SMS and email alerts"
        },
        {
            "num": "Team 3",
            "title": "Module 3: AI-Proctored Examination Suite (EduExam AI)",
            "cat": "Team 3 Workflow & Role Breakdown",
            "summary": "Provides a high-security online examination environment with fullscreen lockdown, client-side AI computer vision proctoring, and network-resilient state sync.",
            "s1": "Student 1 (Frontend & Exam Shell Lead):\n• Fullscreen lockdown browser mode & event traps\n• Disable copy/paste, Alt+Tab, and right-click shortcuts\n• KaTeX formula editor & rich question display\n• Countdown timer with automatic submission trigger",
            "s2": "Student 2 (Backend & Telemetry Architect):\n• Randomized question bank delivery & option shuffler\n• Incremental answer auto-save with idempotency nonces\n• WebSocket heartbeat for disconnection tolerance\n• Automated objective grading & subjective rubric sync",
            "s3": "Student 3 (AI Vision & Integrity Lead):\n• Client MediaPipe face & gaze deviation detection\n• Multi-person presence & missing face alert flags\n• Background audio noise spike analyzer\n• Encrypted violation snapshot aggregator & risk score"
        },
        {
            "num": "Team 4",
            "title": "Module 4: Code Execution Arena & Auto-Judge (EduCode Judge)",
            "cat": "Team 4 Workflow & Role Breakdown",
            "summary": "Scalable coding practice and assessment platform with Monaco IDE, isolated containerized code execution workers, and AST plagiarism detection.",
            "s1": "Student 1 (Frontend & IDE Experience Lead):\n• Monaco Editor integration with multi-language support\n• Custom dark themes, autocomplete, and syntax linting\n• Stdin input console & stdout/stderr execution tabs\n• Real-time problem contest leaderboards & streaks",
            "s2": "Student 2 (Judge Backend & Task Queues):\n• BullMQ distributed task queue in Redis\n• Test case runner evaluating AC, WA, TLE, and MLE\n• Problem creation portal with hidden test cases\n• Execution verdict formatter and runtime benchmark",
            "s3": "Student 3 (Sandbox Security & AST Plagiarism):\n• Docker Engine SDK with gVisor (runsc) kernel jail\n• Strict cgroups limits: 256MB RAM, 1 CPU, no network\n• Python AST tokenization code similarity analyzer\n• Fork bomb mitigation (--pids-limit 64) & tmpfs"
        },
        {
            "num": "Team 5",
            "title": "Module 5: AI Presentation & Virtual Viva (EduPresent AI)",
            "cat": "Team 5 Workflow & Role Breakdown",
            "summary": "Live WebRTC virtual defense studio analyzing vocal delivery, eye contact, and slide alignment, while autonomously generating probing technical viva questions.",
            "s1": "Student 1 (WebRTC Studio & Presentation Lead):\n• Live WebRTC video/audio presentation stage\n• Synchronized PDF/PPTX slide viewer with laser pointer\n• Evaluator interactive rubric scoring panel\n• Audience chat and hand-raise moderation queue",
            "s2": "Student 2 (Multi-Modal AI Pipeline Architect):\n• OpenAI Whisper speech-to-text live transcription\n• Speech pace calculation (words per minute) & pitch\n• Filler-word detector ('um', 'like', 'you know')\n• Real-time AI technical viva question generator",
            "s3": "Student 3 (Computer Vision & Slide Alignment):\n• MediaPipe camera eye contact & posture evaluator\n• Slide text extraction & vector semantic match score\n• DTLS-SRTP WebRTC encryption & slide watermarking\n• Automated composite PDF performance dossier"
        },
        {
            "num": "Team 6",
            "title": "Module 6: RAG Campus Knowledge Assistant (EduAssist AI)",
            "cat": "Team 6 Workflow & Role Breakdown",
            "summary": "Enterprise RAG assistant grounding student and staff queries in verified university handbooks, academic regulations, and syllabus documents.",
            "s1": "Student 1 (Conversational UX Lead):\n• Streaming AI chat UI with markdown and KaTeX\n• Syntax-highlighted code explanation cards\n• Clickable source citation popups (document & page)\n• Student coding tutor mode with guided hints",
            "s2": "Student 2 (RAG Ingestion & Pipeline Lead):\n• PDF/DOCX document text extraction & semantic chunking\n• LangChain / LlamaIndex embedding pipeline\n• Hybrid semantic vector search + keyword BM25\n• Unresolved query ticketing escalation to helpdesk",
            "s3": "Student 3 (Vector Database & Prompt Security):\n• PostgreSQL pgvector extension with HNSW cosine indexes\n• Role-based document access partitions (Student vs Staff)\n• System prompt injection defense & policy filters\n• Latency optimization, Redis caching & query telemetry"
        },
        {
            "num": "Team 7",
            "title": "Module 7: UniHub Git Repo & Research Archive (UniHub)",
            "cat": "Team 7 Workflow & Role Breakdown",
            "summary": "Self-hosted GitHub and arXiv hybrid for university capstone software, Git branch exploration, research paper pre-prints, and campus innovation feeds.",
            "s1": "Student 1 (Git Web UI & Discovery Lead):\n• Git file tree navigator & syntax code browser\n• Visual commit branch graphs & side-by-side diffs\n• Capstone showcase gallery with stars & live demos\n• In-browser PDF research manuscript previewer",
            "s2": "Student 2 (Git Engine & Paper Indexer):\n• Git HTTP smart protocol backend (packfiles, commits)\n• Manuscript metadata parser (title, abstract, authors)\n• Automated BibTeX citation generator & DOI tracking\n• Capstone repository creation wizard with READMEs",
            "s3": "Student 3 (Storage, Activity Feeds & Security):\n• MinIO S3 object store for bare repos & paper PDFs\n• Real-time campus innovation activity event stream\n• Repository access control (Public, Dept, Private)\n• Secret scanning linter preventing API key leaks in code"
        }
    ]

    for t_data in team_slides_data:
        s = prs.slides.add_slide(blank_slide_layout)
        add_header(s, t_data["title"], t_data["cat"])

        # Summary box
        sum_box = add_card(s, Inches(0.8), Inches(1.4), Inches(11.7), Inches(0.7), bg_color=PRIMARY_DARK, border_color=PRIMARY_DARK)
        tb_s = s.shapes.add_textbox(Inches(0.9), Inches(1.45), Inches(11.5), Inches(0.6))
        tf_sum = tb_s.text_frame
        tf_sum.word_wrap = True
        p_s = tf_sum.paragraphs[0]
        p_s.text = t_data["summary"]
        p_s.font.name = "Calibri"
        p_s.font.size = Pt(11)
        p_s.font.color.rgb = WHITE

        # 3 student cards
        roles = [("Student 1: Frontend & UX", t_data["s1"]),
                 ("Student 2: Backend & Logic", t_data["s2"]),
                 ("Student 3: DevOps & Security", t_data["s3"])]
        
        card_w_t = Inches(3.7)
        card_h_t = Inches(4.7)
        start_y_t = Inches(2.3)

        for i, (r_title, r_content) in enumerate(roles):
            cx = Inches(0.8) + i * (card_w_t + Inches(0.3))
            add_card(s, cx, start_y_t, card_w_t, card_h_t)

            tb_r = s.shapes.add_textbox(cx + Inches(0.15), start_y_t + Inches(0.15), card_w_t - Inches(0.3), card_h_t - Inches(0.3))
            tf_r = tb_r.text_frame
            tf_r.word_wrap = True

            p_rt = tf_r.paragraphs[0]
            p_rt.text = r_title
            p_rt.font.name = "Calibri"
            p_rt.font.size = Pt(13)
            p_rt.font.bold = True
            p_rt.font.color.rgb = SECONDARY_BLUE
            p_rt.space_after = Pt(8)

            p_rc = tf_r.add_paragraph()
            p_rc.text = r_content
            p_rc.font.name = "Calibri"
            p_rc.font.size = Pt(10)
            p_rc.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 11: GIT & GITHUB MONOREPO WORKFLOW
    # =========================================================================
    s11 = prs.slides.add_slide(blank_slide_layout)
    add_header(s11, "Git & GitHub Monorepo Collaboration Strategy", "Development Workflow")

    git_cards = [
        ("Turborepo Monorepo Structure", "• apps/web-portal (Central Next.js shell)\n• apps/hrms, apps/attend, apps/exams\n• apps/judge, apps/present, apps/ai, apps/unihub\n• packages/ui (Shared Tailwind/Shadcn tokens)\n• packages/auth (JWT SSO middleware)\n• packages/database (Prisma ORM & PostgreSQL)"),
        ("Branching & PR Protection Rules", "• Main branch: Production locked\n• Staging branch: Integration baseline\n• Branch naming: team<N>/feat-<feature-name>\n• Mandatory 2 approving reviews per PR\n• No direct pushes allowed to main/staging\n• Squash & Merge for linear git histories"),
        ("Automated CI/CD Quality Gates", "• pnpm lint: 0 ESLint warnings required\n• pnpm typecheck: Strict tsc validation\n• pnpm test: Minimum 80% Jest coverage\n• pnpm prisma validate: Schema consistency\n• Docker build verification on pull request\n• Security dependency scan (Trivy / Snyk)")
    ]

    for idx, (g_title, g_content) in enumerate(git_cards):
        cx = Inches(0.8) + idx * (Inches(3.7) + Inches(0.3))
        add_card(s11, cx, Inches(1.5), Inches(3.7), Inches(5.2))
        tb = s11.shapes.add_textbox(cx + Inches(0.2), Inches(1.7), Inches(3.3), Inches(4.8))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = g_title
        p.font.name = "Calibri"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = SECONDARY_BLUE
        p.space_after = Pt(10)

        p2 = tf.add_paragraph()
        p2.text = g_content
        p2.font.name = "Calibri"
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 12: HOW EVERYTHING COMBINES (MASTER INTEGRATION)
    # =========================================================================
    s12 = prs.slides.add_slide(blank_slide_layout)
    add_header(s12, "How Everything Combines: Ingress & Shared Contracts", "Master Integration")

    integ_cards = [
        ("Step 1: Central SSO Gateway", "All 7 modules consume the same JWT authentication claims. A single login grants role-scoped access across HRMS, Attendance, Exams, Code Arena, Presentations, and Repositories."),
        ("Step 2: Shared UI Design System", "Teams build with components from `packages/ui` (buttons, modals, tables, badges). When pages merge into the central Next.js shell, visual design and branding are 100% harmonious."),
        ("Step 3: Database Referential Integrity", "All tables reside in PostgreSQL 16 managed by Prisma ORM. Foreign keys link attendance records to exam eligibility, and student code submissions to capstone repositories."),
        ("Step 4: Cross-Module Event Bus", "Redis Pub/Sub and BullMQ handle cross-module triggers: Attendance <75% locks Exam Hall Tickets; Presentation scores publish to HRMS transcripts; Git commits trigger Activity feed.")
    ]

    for idx, (i_title, i_content) in enumerate(integ_cards):
        cx = Inches(0.8) + (idx % 2) * (Inches(5.7) + Inches(0.3))
        cy = Inches(1.5) + (idx // 2) * (Inches(2.6) + Inches(0.3))
        add_card(s12, cx, cy, Inches(5.7), Inches(2.6))
        tb = s12.shapes.add_textbox(cx + Inches(0.25), cy + Inches(0.2), Inches(5.2), Inches(2.2))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = i_title
        p.font.name = "Calibri"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = SECONDARY_BLUE
        p.space_after = Pt(8)

        p2 = tf.add_paragraph()
        p2.text = i_content
        p2.font.name = "Calibri"
        p2.font.size = Pt(11)
        p2.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 13: TESTING, QA & DEFENSIVE CYBERSECURITY
    # =========================================================================
    s13 = prs.slides.add_slide(blank_slide_layout)
    add_header(s13, "Automated Testing, QA & Defensive Cybersecurity", "Security & Reliability")

    qa_cards = [
        ("OWASP Top 10 Hardening", "• Parameterized queries (Zero SQL injection)\n• Zod request validation on all inputs\n• Argon2id password hashing + TLS 1.3\n• SameSite=Strict cookies + Redis token blacklist\n• Rate limiting on auth & execution APIs"),
        ("Student Code Sandbox Jail", "• gVisor (runsc) kernel-level container jail\n• --network none (Zero outbound network access)\n• Resource limits: 256MB RAM, 1 CPU core\n• Fork bomb termination (--pids-limit 64)\n• Ephemeral in-memory tmpfs (/tmp:rw,16m)"),
        ("Anti-Cheat & Anti-Proxy", "• 5s rotating HMAC QR eliminates screenshots\n• Geofence polygon + mock GPS detector\n• MediaPipe face liveness & gaze deviation flags\n• WebRTC DTLS-SRTP encrypted streams\n• Watermarked capstone presentation slides")
    ]

    for idx, (q_title, q_content) in enumerate(qa_cards):
        cx = Inches(0.8) + idx * (Inches(3.7) + Inches(0.3))
        add_card(s13, cx, Inches(1.5), Inches(3.7), Inches(5.2))
        tb = s13.shapes.add_textbox(cx + Inches(0.2), Inches(1.7), Inches(3.3), Inches(4.8))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = q_title
        p.font.name = "Calibri"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = SECONDARY_BLUE
        p.space_after = Pt(10)

        p2 = tf.add_paragraph()
        p2.text = q_content
        p2.font.name = "Calibri"
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 14: FINAL STAGING INTEGRATION & EVALUATION RUBRIC
    # =========================================================================
    s14 = prs.slides.add_slide(blank_slide_layout)
    add_header(s14, "Staging Integration & Capstone Evaluation Rubric", "Evaluation & Assessment")

    eval_items = [
        ("Architecture & Code Quality (25%)", "Adherence to clean architecture, TypeScript strict typing, normalized database schema, zero lint warnings, and modular folder structure."),
        ("Feature Completeness & Zero Broken Flows (30%)", "100% functional user journeys. Every button leads to an active page; no placeholders, no 'to-be-implemented' comments, full error boundaries."),
        ("Multi-Modal AI & Security Hardening (25%)", "Robust implementation of AI models (face liveness, speech transcription, vector RAG, code sandbox jail) and verified anti-cheating mechanisms."),
        ("Team Collaboration & Git Hygiene (20%)", "Active participation by all 3 team members, adherence to PR review standards, clean commit messages, and automated CI passing.")
    ]

    for idx, (e_title, e_desc) in enumerate(eval_items):
        cy = Inches(1.5) + idx * Inches(1.35)
        add_card(s14, Inches(0.8), cy, Inches(11.7), Inches(1.15))
        tb = s14.shapes.add_textbox(Inches(1.0), cy + Inches(0.12), Inches(11.3), Inches(0.9))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = e_title
        p.font.name = "Calibri"
        p.font.size = Pt(13)
        p.font.bold = True
        p.font.color.rgb = SECONDARY_BLUE

        p2 = tf.add_paragraph()
        p2.text = e_desc
        p2.font.name = "Calibri"
        p2.font.size = Pt(10.5)
        p2.font.color.rgb = TEXT_DARK

    # =========================================================================
    # SLIDE 15: SUMMARY & NEXT STEPS FOR TEAMS
    # =========================================================================
    s15 = prs.slides.add_slide(blank_slide_layout)
    bg15 = s15.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg15.fill.solid()
    bg15.fill.fore_color.rgb = PRIMARY_DARK
    bg15.line.fill.background()

    tb_end = s15.shapes.add_textbox(Inches(1.2), Inches(1.2), Inches(11.0), Inches(5.2))
    tf_end = tb_end.text_frame
    tf_end.word_wrap = True

    p_e0 = tf_end.paragraphs[0]
    p_e0.text = "READY FOR EXECUTION: NEXT ACTIONS FOR STUDENT TEAMS"
    p_e0.font.name = "Calibri"
    p_e0.font.size = Pt(14)
    p_e0.font.bold = True
    p_e0.font.color.rgb = ACCENT_TEAL
    p_e0.space_after = Pt(10)

    p_e1 = tf_end.add_paragraph()
    p_e1.text = "UniSphere Master Implementation Kickoff"
    p_e1.font.name = "Calibri"
    p_e1.font.size = Pt(26)
    p_e1.font.bold = True
    p_e1.font.color.rgb = WHITE
    p_e1.space_after = Pt(18)

    action_points = [
        "1. Team Formation: 7 teams of 3 students finalize role assignments (Student 1: Frontend, Student 2: Backend, Student 3: DevOps/Security).",
        "2. Monorepo Clone: Clone the repository, run `pnpm install`, and verify local `docker-compose up` launches PostgreSQL, Redis, and MinIO.",
        "3. Sprint 1 Kickoff: Implement core database migrations from `packages/database` and establish basic API endpoints.",
        "4. Interface Contract Freeze: Review OpenAPI specifications in `PROJECT_SPECIFICATION.md` to ensure zero cross-module breaking changes.",
        "5. Weekly Standups: Conduct weekly bi-directional code reviews and demo working prototypes at the end of each two-week sprint."
    ]
    for pt in action_points:
        p_pt = tf_end.add_paragraph()
        p_pt.text = pt
        p_pt.font.name = "Calibri"
        p_pt.font.size = Pt(12)
        p_pt.font.color.rgb = RGBColor(0xCB, 0xD5, 0xE1)
        p_pt.space_after = Pt(6)

    # Save presentation
    base_dir = "c:/Users/bdeepak/OneDrive - Chanakya University/Course materials/projects"
    ppt_path = os.path.join(base_dir, "UNISPHERE_PROJECT_ROADMAP.pptx")
    prs.save(ppt_path)
    print(f"Successfully generated PowerPoint presentation at {ppt_path}")

if __name__ == "__main__":
    create_presentation()
